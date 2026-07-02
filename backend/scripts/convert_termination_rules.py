from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE_CANDIDATES = [
    ROOT / "终止登记审查规则_分类版.xlsx",
    ROOT / "审查规则" / "终止登记审查规则_分类版.xlsx",
]
SOURCE = next((path for path in SOURCE_CANDIDATES if path.exists()), SOURCE_CANDIDATES[0])
OUTPUT = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "rules"
    / "rules_termination.json"
)

TERMINATION_FIELD_ANCHORS = [
    "清算报告日期",
    "实际到期日期",
    "是否按约定日期清算",
    "是否已完成信托财产分配",
    "是否已完成销户",
    "累计实收信托",
    "日均实收信托",
    "实际收益",
    "信托费用",
    "受托人累计基本报酬",
    "受托人累计业绩报酬",
    "信托收益累计分配额",
    "信托本金累计给付额",
    "赔付金额",
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_dimension(value: str) -> str:
    raw = clean(value)
    if not raw:
        return "其他"
    raw = re.sub(r"^\d+\s*[.．、]\s*", "", raw).strip()
    mapping = {
        "格式模板": "格式模板规则库",
        "登记必填要素": "登记必填要素规则库",
        "跨材料数据逻辑": "跨材料数据逻辑校验库",
        "监管合规红线": "监管合规红线规则库",
        "业务退回/整改": "业务退回/整改案例库",
        "审查风险分级": "审查风险分级规则库",
    }
    return mapping.get(raw, raw)


def check_type(rule_name: str, rule_text: str, dimension: str) -> str:
    text = f"{rule_name} {rule_text} {dimension}"
    if any(kw in text for kw in ("一致", "相一致", "相符", "对应")):
        return "跨材料一致性"
    if any(kw in text for kw in ("必须包含", "应上传", "未记载", "佐证材料", "材料提交", "申请材料")):
        return "材料完整性"
    if any(kw in text for kw in ("责任解除", "监管", "合规", "受托人", "赔付", "违反信托目的")):
        return "监管口径符合性"
    return "语义条件判断"


def risk_level(rule_name: str, rule_text: str, dimension: str) -> str:
    text = f"{rule_name} {rule_text} {dimension}"
    if any(kw in text for kw in ("必须包含", "应上传", "缺失", "未提供", "一致性", "相一致", "清算报告与终止信息")):
        return "高风险"
    if "跨材料" in dimension or "监管" in dimension:
        return "中风险"
    return "中风险"


def basis_file(source_file: str, basis_text: str) -> str:
    src = clean(source_file)
    if src:
        return src
    text = clean(basis_text)
    if not text:
        return ""
    head = re.split(r"[：:；;\n]", text, maxsplit=1)[0].strip()
    return head[:80]


def file_applicable_materials(file_name: str, rule_name: str, rule_text: str) -> list[str]:
    text = f"{file_name} {rule_name} {rule_text}"
    if "清算报告" in text and not any(kw in text for kw in ("一致", "终止信息", "对应")):
        return ["清算报告", "其他附件"]
    if any(kw in text for kw in ("一致", "终止信息", "佐证")):
        return ["终止登记申报模板JSON", "清算报告", "其他附件"]
    if "申请书" in text:
        return ["终止登记申请书", "其他附件"]
    return ["终止登记申报模板JSON", "终止登记申请书", "清算报告", "其他附件"]


def file_machine_params(rule_id: str, rule_text: str) -> str:
    if rule_id in {"TERM-FILE-AI-001", "TERM-FILE-AI-004"}:
        return (
            "rule_role=overall_consistency_fallback; "
            f"covered_by_field_rules={','.join(TERMINATION_FIELD_ANCHORS)}; "
            "do_not_duplicate_field_rules=true; "
            "output_only_unanchored_or_material_wide_inconsistency=true"
        )
    if rule_id == "TERM-FILE-AI-002":
        return "rule_role=material_required; required_material=受托人出具的清算报告"
    if rule_id == "TERM-FILE-AI-003":
        return (
            "rule_role=responsibility_release_condition; field_anchor=责任解除条件; "
            "check_property_distributed=true; check_account_closed=true; "
            "check_beneficiary_no_objection=true"
        )
    if rule_id == "TERM-FILE-AI-005":
        return (
            "rule_role=supporting_evidence_fallback; "
            "only_when_liquidation_report_lacks_corresponding_content_and_no_supporting_attachment=true; "
            "do_not_report_if_report_records_unfinished_status=true"
        )
    return ""


def file_ai_focus(rule_id: str) -> list[str]:
    if rule_id in {"TERM-FILE-AI-001", "TERM-FILE-AI-004"}:
        return [
            "仅兜底核对未被具体要素规则覆盖的材料层一致性问题",
            "如同一事实已落入清算报告日期、实际到期日期、财务金额、财产分配、销户、赔付等具体字段规则，应使用具体要素规则输出",
            "指出支撑判断的材料位置和文本片段",
        ]
    if rule_id == "TERM-FILE-AI-002":
        return [
            "判断是否已上传受托人出具的清算报告",
            "清算报告已存在时不得输出材料缺失问题",
            "材料缺失时 value 必须写成 缺失:受托人出具的清算报告",
        ]
    if rule_id == "TERM-FILE-AI-003":
        return [
            "核对受托人责任解除条件是否满足",
            "重点检查财产是否已全部分配、信托财产专户是否已销户、受益人是否无异议或已有有效说明",
            "不要与具体字段一致性问题重复输出",
        ]
    if rule_id == "TERM-FILE-AI-005":
        return [
            "仅当清算报告未记载终止登记对应内容且没有其他佐证附件时输出",
            "清算报告已记载未完成分配、未销户等状态时，不得表述为未记载",
            "指出缺少的佐证内容或附件名称",
        ]
    return [
        "判断终止登记申请材料是否齐备",
        "核对终止登记信息与清算报告或佐证附件是否一致",
        "指出支撑判断的材料位置和文本片段",
    ]


def inherited_rows(ws, headers: list[str], fill_columns: set[str]) -> list[tuple[int, dict[str, str]]]:
    last: dict[str, str] = {}
    rows: list[tuple[int, dict[str, str]]] = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        current: dict[str, str] = {}
        for idx, header in enumerate(headers):
            value = clean(row[idx]) if idx < len(row) else ""
            if not value and header in fill_columns:
                value = last.get(header, "")
            current[header] = value
            if value:
                last[header] = value
        rows.append((row_no, current))
    return rows


def file_review_rules(ws) -> list[dict[str, Any]]:
    headers = [clean(c.value) for c in ws[1]]
    fill_columns = {
        "登记流程",
        "文件名",
        "文件审查规则",
        "审核维度",
        "具体的规则",
        "审核依据",
        "来源文件",
    }
    rules: list[dict[str, Any]] = []
    for _, (row_no, row) in enumerate(inherited_rows(ws, headers, fill_columns), start=1):
        if row.get("分类", "").lower() != "ai":
            continue
        rule_text = row.get("具体的规则", "")
        if not rule_text:
            continue
        rule_name = row.get("文件审查规则", "") or row.get("文件名", "") or "终止登记文件审查"
        dim = normalize_dimension(row.get("审核维度", ""))
        basis_text = row.get("审核依据", "")
        idx = len(rules) + 1
        rule_id = f"TERM-FILE-AI-{idx:03d}"
        rules.append(
            {
                "rule_id": rule_id,
                "registration_type": "终止登记",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": row_no,
                "rule_name": rule_name,
                "rule_text": rule_text,
                "basis_file": basis_file(row.get("来源文件", ""), basis_text),
                "basis_text": basis_text,
                "review_dimension": dim,
                "table_name": "",
                "field_name": "",
                "applicable_materials": file_applicable_materials(
                    row.get("文件名", ""), rule_name, rule_text
                ),
                "review_method": "ai",
                "original_review_method": "AI",
                "check_type": check_type(rule_name, rule_text, dim),
                "trigger_condition": row.get("文件名", ""),
                "machine_params": file_machine_params(rule_id, rule_text),
                "risk_level": risk_level(rule_name, rule_text, dim),
                "ai_check_focus": file_ai_focus(rule_id),
                "evidence_requirement": "请 AI 在审核时指出涉及的材料、表名、字段、页码或文本片段。",
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def element_rules(ws) -> list[dict[str, Any]]:
    headers = [clean(c.value) for c in ws[1]]
    fill_columns = {
        "登记流程",
        "文件名",
        "表名",
        "要素名",
        "审查维度",
        "具体规则",
        "审核依据（法规或者文件里规定怎么写）",
    }
    rules: list[dict[str, Any]] = []
    for _, (row_no, row) in enumerate(inherited_rows(ws, headers, fill_columns), start=1):
        if row.get("分类", "").lower() != "ai":
            continue
        rule_text = row.get("具体规则", "")
        if not rule_text:
            continue
        field_name = row.get("要素名", "") or "终止登记要素审查"
        table_name = row.get("表名", "")
        dim = normalize_dimension(row.get("审查维度", ""))
        basis_text = row.get("审核依据（法规或者文件里规定怎么写）", "")
        idx = len(rules) + 1
        rules.append(
            {
                "rule_id": f"TERM-ELEMENT-AI-{idx:03d}",
                "registration_type": "终止登记",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": row_no,
                "rule_name": field_name,
                "rule_text": rule_text,
                "basis_file": basis_file("", basis_text),
                "basis_text": basis_text,
                "review_dimension": dim,
                "table_name": table_name,
                "field_name": field_name,
                "applicable_materials": ["终止登记申报模板JSON", "清算报告", "其他附件"],
                "review_method": "ai",
                "original_review_method": "AI",
                "check_type": check_type(field_name, rule_text, dim),
                "trigger_condition": "",
                "machine_params": (
                    f"rule_role=field_primary; field_anchor={field_name}; "
                    "issue_grain=one_field_fact; prefer_specific_rule=true"
                ),
                "risk_level": risk_level(field_name, rule_text, dim),
                "ai_check_focus": [
                    "优先读取终止登记申报模板中的表名和字段值",
                    "结合清算报告、申请书或其他附件判断字段值是否符合填报口径",
                    "同一事实不要再用 TERM-FILE-AI-001 或 TERM-FILE-AI-004 重复输出",
                    "输出问题时定位到具体表名、字段、页码或文本片段",
                ],
                "evidence_requirement": "请 AI 在审核时指出涉及的材料、表名、字段、页码或文本片段。",
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    rules = file_review_rules(wb["文件审查"]) + element_rules(wb["要素审查"])
    payload = {
        "rule_set_id": "termination_ai_rules_v1",
        "registration_type": "终止登记",
        "source_file": SOURCE.name,
        "description": "终止登记 AI 审核规则；来源于《终止登记审查规则_分类版.xlsx》中分类标记为 ai 的文件审查和要素审查规则。",
        "total_rules": len(rules),
        "rules": rules,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {OUTPUT} ({len(rules)} rules)")


if __name__ == "__main__":
    main()
