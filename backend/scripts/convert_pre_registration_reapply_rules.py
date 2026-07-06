from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "审查规则" / "重新申请预登记审查规则表.xlsx"
OUTPUT = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "rules"
    / "rules_pre_registration_reapply.json"
)

FILE_SHEET = "文件审查规则"
ELEMENT_SHEET = "要素审查"
ELEMENT_AI_MARK = "需借助AI推理判断"


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_dimension(value: str) -> str:
    raw = clean(value)
    if not raw:
        return "其他"
    if raw.endswith("规则库") or raw.endswith("案例库"):
        return raw
    mapping = {
        "格式模板": "格式模板规则库",
        "登记必填要素": "登记必填要素规则库",
        "跨材料数据逻辑": "跨材料数据逻辑校验库",
        "监管合规红线": "监管合规红线规则库",
        "业务退回/整改": "业务退回/整改案例库",
        "审查风险分级": "审查风险分级规则库",
    }
    return mapping.get(raw, raw)


def split_basis(raw: str) -> tuple[str, str]:
    text = clean(raw)
    if not text:
        return "", ""
    head = (
        text.split("—", 1)[0]
        .split("：", 1)[0]
        .split("“", 1)[0]
        .strip(" ，。；;")
    )
    return head[:80], text


def normalize_table_name(value: str) -> str:
    raw = clean(value)
    match = re.match(r"^\d+[.．]\s*(.+)$", raw)
    return match.group(1).strip() if match else raw


def field_anchor(table_name: str, field_name: str) -> str:
    return field_name or table_name


def severity_type(dimension: str) -> str:
    return "risk_hint" if "审查风险分级" in dimension else "issue"


def risk_level(rule_text: str, dimension: str) -> str:
    text = clean(rule_text)
    if "审查风险分级" in dimension:
        return "低风险"
    if any(kw in text for kw in ("不得", "严禁", "红线", "违法", "违规", "不涉及新增地方政府融资平台债务")):
        return "高风险"
    if any(kw in text for kw in ("缺失", "未提交", "未上传", "应提交", "应上传", "必须", "必传", "不一致", "矛盾")):
        return "高风险"
    if "监管合规" in dimension or "跨材料" in dimension or "业务退回" in dimension:
        return "中风险"
    return "中风险"


def check_type(rule_text: str, dimension: str) -> str:
    text = clean(rule_text)
    if "签字" in text or "盖章" in text or "公章" in text or "CA认证" in text:
        return "签字盖章存在性"
    if "一致" in text or "相符" in text or "同一" in text:
        return "跨材料一致性"
    if any(kw in text for kw in ("上传", "提交", "材料", "文件", "附件", "必传")):
        return "材料完整性"
    if "监管" in dimension or "合规" in dimension:
        return "监管口径符合性"
    return "语义条件判断"


def issue_type(rule_text: str, dimension: str, source_sheet: str) -> str:
    text = clean(rule_text)
    if "审查风险分级" in dimension:
        return "风险提示"
    if source_sheet == FILE_SHEET and any(kw in text for kw in ("上传", "提交", "材料", "文件", "附件", "必传")):
        return "材料必交"
    if "签字" in text or "盖章" in text or "公章" in text or "CA认证" in text:
        return "签字盖章"
    if "一致" in text or "相符" in text or "同一" in text:
        return "跨材料一致性"
    if "监管合规" in dimension or "合规" in text or "不得" in text:
        return "监管红线"
    if "必填" in text or "不得为空" in text or "只能填写" in text:
        return "条件必填"
    return "语义判断"


def rule_role(rule_text: str, dimension: str, source_sheet: str) -> str:
    itype = issue_type(rule_text, dimension, source_sheet)
    return {
        "材料必交": "material_required",
        "签字盖章": "signature_or_seal",
        "跨材料一致性": "cross_document_consistency",
        "监管红线": "regulatory_redline",
        "条件必填": "conditional_required",
        "风险提示": "risk_hint",
    }.get(itype, "semantic_check")


def applicable_materials_for_file(file_name: str, rule_name: str, rule_text: str) -> list[str]:
    text = f"{file_name} {rule_name} {rule_text}"
    materials: list[str] = []
    if "申请书" in text:
        materials.append("申请书")
    if "模板" in text or "JSON" in text or "json" in text:
        materials.append("申报模板")
    if "要素报告表" in text:
        materials.append("信托预登记要素报告表")
    if "政信" in text or "融资平台债务" in text:
        materials.append("政信类证明材料")
    if "新型资产服务信托" in text or "情况说明" in text:
        materials.append("新型资产服务信托情况说明")
    if "其他附件" in text or "其他文件" in text:
        materials.append("其他附件")
    if not materials:
        materials = ["申请书", "申报模板", "其他附件"]
    return list(dict.fromkeys(materials))


def applicable_materials_for_element(rule_text: str) -> list[str]:
    text = clean(rule_text)
    materials = ["申报模板"]
    if any(kw in text for kw in ("申请书", "合规承诺书", "签字", "盖章", "公章")):
        materials.append("申请书")
    if any(kw in text for kw in ("信托合同", "产品说明书", "备案表单", "证明材料", "情况说明", "附件")):
        materials.append("其他附件")
    return materials


def skip_when(rule_text: str, table_name: str, field_name: str, source_sheet: str) -> str:
    text = clean(rule_text)
    if source_sheet == FILE_SHEET:
        if "要素报告表PDF" in text and "未上传要素报告表PDF时" in text:
            return "信托预登记要素报告表PDF缺失"
        if "未通过CA认证" in text or "使用电子认证服务" in text:
            return "CA认证登录状态 unknown"
        if "政信类业务" in text or "地方政府融资平台债务" in text:
            return "未识别到政信类业务特征"
        if "新型资产服务信托" in text:
            return "资产服务信托分类2 != 新型资产服务信托"
        return ""
    enum_check = any(kw in text for kw in ("枚举范围", "有效选项", "可选值"))
    if field_name == "交易对手是否隐债主体" and enum_check:
        return ""
    if field_name == "保（托）管人名称":
        return "托管信息.是否聘请保（托）管人 != 是"
    if field_name == "重大关联交易是否已经董事会批准":
        return "关联交易信息 empty OR 关联交易性质 != 重大关联交易"
    if table_name == "关联交易信息" or field_name.startswith("关联交易"):
        return "关联交易信息 empty"
    if (
        "政信" in text
        or "隐债" in text
        or "地方政府融资平台" in text
        or (field_name == "交易对手是否隐债主体" and not enum_check)
    ):
        return "底层资产及交易对手.交易对手是否隐债主体 != 是"
    if "房地产" in text or field_name in {"房地产项目类型", "房地产项目类型详情", "房地产开发企业名称", "资本金比例", "开发商资质"}:
        return "房地产项目信息 empty"
    if "异地推介" in text or field_name.startswith("推介"):
        return "异地推介信息 empty"
    return ""


def ai_focus(rule_text: str, source_sheet: str) -> list[str]:
    text = clean(rule_text)
    if source_sheet == FILE_SHEET:
        focus = [
            "判断重新申请预登记申请材料是否齐备、形式是否符合要求",
            "结合申请书、申报模板 JSON、要素报告表 PDF 和其他附件识别缺失、不一致或证明不充分的问题",
            "输出问题时定位到具体材料、页码、字段或文本片段",
        ]
        if "Excel原文件" in text or "JSON" in text:
            focus.append("本网页审核以上传导出的预登记产品EXCEL模板.json为准，不要求上传 Excel 原文件")
        return focus
    return [
        "优先在预登记申报模板 JSON 的对应表和字段中判断规则是否命中",
        "如规则涉及申请书、签章、证明材料或附件，再结合相关材料判断",
        "输出问题时定位到具体表名、记录和字段",
    ]


def row_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, str]:
    return {header: clean(row[idx]) if idx < len(row) else "" for idx, header in enumerate(headers)}


def file_review_rules(ws) -> list[dict[str, Any]]:
    headers = [clean(c.value) for c in ws[1]]
    rules: list[dict[str, Any]] = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw = row_dict(headers, row)
        if raw.get("分类", "").lower() != "ai":
            continue
        file_name = raw["文件名"]
        rule_name = raw["文件审查规则"]
        rule_text = raw["具体的规则"]
        if not rule_text:
            continue
        dim = normalize_dimension(raw["审核维度"])
        basis_file, basis_text = split_basis(raw["审核依据"])
        rid = f"REREG-FILE-AI-{len(rules) + 1:03d}"
        rules.append(
            {
                "rule_id": rid,
                "registration_type": "重新申请预登记",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": row_no,
                "rule_name": rule_name or file_name or "重新申请预登记文件审查",
                "rule_text": rule_text,
                "basis_file": basis_file,
                "basis_text": basis_text,
                "review_dimension": dim,
                "table_name": "",
                "field_name": "",
                "applicable_materials": applicable_materials_for_file(file_name, rule_name, rule_text),
                "review_method": "ai",
                "original_review_method": "AI",
                "check_type": check_type(rule_text, dim),
                "trigger_condition": file_name,
                "machine_params": "",
                "risk_level": risk_level(rule_text, dim),
                "ai_check_focus": ai_focus(rule_text, ws.title),
                "evidence_requirement": "请 AI 指出涉及的材料、表名、字段、页码或文本片段。",
                "rule_role": rule_role(rule_text, dim, ws.title),
                "severity_type": severity_type(dim),
                "issue_type": issue_type(rule_text, dim, ws.title),
                "field_anchor": "",
                "skip_when": skip_when(rule_text, "", "", ws.title),
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def element_ai_rules(ws) -> list[dict[str, Any]]:
    headers = [clean(c.value) for c in ws[1]]
    rules: list[dict[str, Any]] = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        raw = row_dict(headers, row)
        if raw.get("分类结果") != ELEMENT_AI_MARK:
            continue
        table_name = normalize_table_name(raw["要素分类"])
        field_name = clean(raw["要素名"])
        if field_name == raw["要素分类"]:
            field_name = ""
        rule_text = raw["具体规则(已融合全部枚举值)"]
        if not rule_text:
            continue
        dim = normalize_dimension(raw["审查维度"])
        basis_file, basis_text = split_basis(raw["审查依据"])
        rid = f"REREG-ELEMENT-AI-{len(rules) + 1:03d}"
        rules.append(
            {
                "rule_id": rid,
                "registration_type": "重新申请预登记",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": row_no,
                "rule_name": field_name or table_name or "重新申请预登记要素审查",
                "rule_text": rule_text,
                "basis_file": basis_file,
                "basis_text": basis_text,
                "review_dimension": dim,
                "table_name": table_name,
                "field_name": field_name,
                "applicable_materials": applicable_materials_for_element(rule_text),
                "review_method": "ai",
                "original_review_method": ELEMENT_AI_MARK,
                "check_type": check_type(rule_text, dim),
                "trigger_condition": "",
                "machine_params": "",
                "risk_level": risk_level(rule_text, dim),
                "ai_check_focus": ai_focus(rule_text, ws.title),
                "evidence_requirement": "请 AI 在审核时指出涉及的材料、表名、字段、页码或文本片段。",
                "rule_role": rule_role(rule_text, dim, ws.title),
                "severity_type": severity_type(dim),
                "issue_type": issue_type(rule_text, dim, ws.title),
                "field_anchor": field_anchor(table_name, field_name),
                "skip_when": skip_when(rule_text, table_name, field_name, ws.title),
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def validate_rules(rules: list[dict[str, Any]]) -> None:
    ids = [rule["rule_id"] for rule in rules]
    if len(ids) != len(set(ids)):
        raise RuntimeError("生成规则 ID 重复")
    required = {
        "rule_id",
        "registration_type",
        "rule_name",
        "rule_text",
        "review_dimension",
        "applicable_materials",
        "risk_level",
        "enabled",
        "demo_enabled",
    }
    for rule in rules:
        missing = [key for key in required if key not in rule or rule[key] in ("", [])]
        if missing:
            raise RuntimeError(f"{rule.get('rule_id')} 缺少字段: {missing}")


def main() -> None:
    wb = load_workbook(SOURCE, data_only=True, read_only=True)
    rules = file_review_rules(wb[FILE_SHEET]) + element_ai_rules(wb[ELEMENT_SHEET])
    validate_rules(rules)
    payload = {
        "rule_set_id": "pre_registration_reapply_ai_rules_v2",
        "registration_type": "重新申请预登记",
        "source_file": str(SOURCE.relative_to(ROOT)),
        "description": "重新申请预登记 AI 审核内置规则；来源于审查规则文件夹《重新申请预登记审查规则表.xlsx》中，文件审查规则 sheet 标记为 ai 的规则，以及要素审查 sheet 分类结果为需借助AI推理判断的规则。",
        "total_rules": len(rules),
        "rules": rules,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        f"wrote {OUTPUT} ({len(rules)} rules: "
        f"{sum(1 for r in rules if r['source_sheet'] == FILE_SHEET)} file, "
        f"{sum(1 for r in rules if r['source_sheet'] == ELEMENT_SHEET)} element)"
    )


if __name__ == "__main__":
    main()
