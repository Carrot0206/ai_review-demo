from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE_CANDIDATES = [
    ROOT / "预登记审查规则_分类结果.xlsx",
    ROOT / "审查规则" / "预登记审查规则_分类结果.xlsx",
]
SOURCE = next((p for p in SOURCE_CANDIDATES if p.exists()), SOURCE_CANDIDATES[0])
OUTPUT = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "rules"
    / "rules_pre_registration.json"
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_dimension(value: str) -> str:
    if not value:
        return "其他"
    if value.endswith("规则库") or value.endswith("案例库"):
        return value
    mapping = {
        "格式模板": "格式模板规则库",
        "登记必填要素": "登记必填要素规则库",
        "跨材料数据逻辑": "跨材料数据逻辑校验库",
        "监管合规红线": "监管合规红线规则库",
        "业务退回/整改": "业务退回/整改案例库",
        "审查风险分级": "审查风险分级规则库",
    }
    return mapping.get(value, value)


def risk_level(rule_text: str, dimension: str) -> str:
    text = rule_text or ""
    if any(kw in text for kw in ("未提供", "请提供", "缺少", "需上传", "均需上传", "必须", "不得为空")):
        return "高风险"
    if "监管合规" in dimension or "业务退回" in dimension or "跨材料" in dimension:
        return "中风险"
    return "中风险"


def severity_type(dimension: str) -> str:
    return "risk_hint" if "审查风险分级" in dimension else "issue"


def issue_type(rule_text: str, dimension: str, source_sheet: str) -> str:
    text = rule_text or ""
    if "审查风险分级" in dimension:
        return "风险提示"
    if source_sheet == "文件审查":
        return "材料必交" if any(kw in text for kw in ("上传", "提交", "提供")) else "文件审查"
    if "缺少" in text or "请提供" in text or "上传" in text or "提交" in text:
        return "材料必交"
    if "不得空置" in text or "不得为空" in text or "须" in text or "必须" in text:
        return "条件必填"
    if "一致" in text:
        return "跨字段一致性"
    if "监管合规" in dimension:
        return "监管红线"
    if "标准化意见" in text:
        return "标准化意见案例"
    return "语义判断"


def rule_role(rule_text: str, dimension: str, source_sheet: str) -> str:
    if source_sheet == "文件审查":
        return "material_required" if issue_type(rule_text, dimension, source_sheet) == "材料必交" else "file_review"
    if severity_type(dimension) == "risk_hint":
        return "risk_hint"
    return {
        "材料必交": "material_required",
        "条件必填": "conditional_required",
        "跨字段一致性": "cross_field_consistency",
        "监管红线": "regulatory_redline",
        "标准化意见案例": "standard_opinion_case",
    }.get(issue_type(rule_text, dimension, source_sheet), "semantic_check")


def field_anchor(table_name: str, field_name: str) -> str:
    return field_name or table_name


def skip_when(rule_id: str, table_name: str, field_name: str, rule_text: str) -> str:
    enum_check = any(kw in rule_text for kw in ("枚举范围", "有效选项", "可选值"))
    if field_name == "交易对手是否隐债主体" and enum_check:
        return ""
    if field_name == "保（托）管人名称":
        return "托管信息.是否聘请保（托）管人 != 是"
    if field_name == "重大关联交易是否已经董事会批准":
        return "关联交易信息 empty OR 关联交易性质 != 重大关联交易"
    if table_name == "关联交易事项" or field_name.startswith("关联交易"):
        return "关联交易信息 empty"
    if (
        "政信" in rule_text
        or "隐债" in rule_text
        or "地方政府融资平台" in rule_text
        or (field_name == "交易对手是否隐债主体" and not enum_check)
    ):
        return "底层资产及交易对手.交易对手是否隐债主体 != 是"
    if "房地产" in rule_text or field_name in {"房地产项目类型", "房地产项目类型详情", "房地产开发企业名称", "资本金比例", "开发商资质"}:
        return "房地产项目信息 empty"
    if "异地推介" in rule_text or field_name.startswith("推介"):
        return "异地推介信息 empty"
    return ""


def check_type(rule_text: str, dimension: str) -> str:
    text = rule_text or ""
    if "盖章" in text or "签字" in text or "公章" in text:
        return "签字盖章存在性"
    if "上传" in text or "提供" in text or "材料" in text or "文件" in text:
        return "材料完整性"
    if "一致" in text or "相符" in text:
        return "跨材料一致性"
    if "监管" in dimension or "合规" in dimension:
        return "监管口径符合性"
    return "语义条件判断"


def split_basis(raw: str) -> tuple[str, str]:
    if not raw:
        return "", ""
    head = raw.split("—", 1)[0].split("：", 1)[0].strip()
    return head[:80], raw


def normalize_section_or_field(value: str) -> tuple[str, str]:
    raw = clean(value)
    if not raw:
        return "", ""
    match = re.match(r"^\d+[.．]\s*(.+)$", raw)
    if match:
        return match.group(1).strip(), ""
    return "", raw


def file_review_rules(ws) -> list[dict[str, Any]]:
    rules: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        file_item, process, rule_name, dimension, rule_text, basis, source_file = [clean(v) for v in row[:7]]
        if not rule_text:
            continue
        basis_file, basis_text = split_basis(basis)
        dim = normalize_dimension(dimension)
        itype = issue_type(rule_text, dim, ws.title)
        rules.append(
            {
                "rule_id": f"PREG-FILE-AI-{idx:03d}",
                "registration_type": "预登记",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": idx + 1,
                "rule_name": rule_name or file_item or "预登记文件审查",
                "rule_text": rule_text,
                "basis_file": basis_file or source_file,
                "basis_text": basis_text or basis,
                "review_dimension": dim,
                "table_name": "",
                "field_name": "",
                "applicable_materials": ["申请书", "申报模板", "其他附件"],
                "review_method": "ai",
                "original_review_method": "AI",
                "check_type": check_type(rule_text, dim),
                "trigger_condition": process,
                "machine_params": "",
                "risk_level": risk_level(rule_text, dim),
                "ai_check_focus": [
                    "判断预登记申请材料是否齐备",
                    "指出缺失或不符合要求的文件名称",
                    "说明缺失问题是否影响审核结论",
                ],
                "evidence_requirement": "请 AI 在审核时指出涉及的材料、表名、字段、页码或文本片段。",
                "rule_role": rule_role(rule_text, dim, ws.title),
                "severity_type": severity_type(dim),
                "issue_type": itype,
                "field_anchor": field_anchor("", ""),
                "skip_when": skip_when(f"PREG-FILE-AI-{idx:03d}", "", "", rule_text),
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def element_ai_rules(ws) -> list[dict[str, Any]]:
    rules: list[dict[str, Any]] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        _, dimension, element_name, rule_text, basis = [clean(v) for v in row[:5]]
        if not rule_text:
            continue
        table_name, field_name = normalize_section_or_field(element_name)
        basis_file, basis_text = split_basis(basis)
        dim = normalize_dimension(dimension)
        rid = f"PREG-ELEMENT-AI-{idx:03d}"
        itype = issue_type(rule_text, dim, ws.title)
        rules.append(
            {
                "rule_id": rid,
                "registration_type": "预登记",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": idx + 1,
                "rule_name": field_name or table_name or "预登记要素审查",
                "rule_text": rule_text,
                "basis_file": basis_file,
                "basis_text": basis_text or basis,
                "review_dimension": dim,
                "table_name": table_name,
                "field_name": field_name,
                "applicable_materials": ["申报模板", "申请书", "其他附件"],
                "review_method": "ai",
                "original_review_method": "AI",
                "check_type": check_type(rule_text, dim),
                "trigger_condition": "",
                "machine_params": "",
                "risk_level": risk_level(rule_text, dim),
                "ai_check_focus": [
                    "优先在预登记申报模板字段中判断规则是否命中",
                    "如规则涉及文件、签章或承诺材料，再结合申请书或附件判断",
                    "输出问题时定位到具体表名、记录和字段",
                ],
                "evidence_requirement": "请 AI 在审核时指出涉及的材料、表名、字段、页码或文本片段。",
                "rule_role": rule_role(rule_text, dim, ws.title),
                "severity_type": severity_type(dim),
                "issue_type": itype,
                "field_anchor": field_anchor(table_name, field_name),
                "skip_when": skip_when(rid, table_name, field_name, rule_text),
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    rules = file_review_rules(wb["文件审查"]) + element_ai_rules(wb["需AI推理判断规则"])
    payload = {
        "rule_set_id": "pre_registration_ai_rules_v1",
        "registration_type": "预登记",
        "source_file": SOURCE.name,
        "description": "预登记 AI 审核规则；来源于《预登记审查规则_分类结果.xlsx》的文件审查和需AI推理判断规则两个 sheet。",
        "total_rules": len(rules),
        "rules": rules,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {OUTPUT} ({len(rules)} rules)")


if __name__ == "__main__":
    main()
