from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE_CANDIDATES = [
    ROOT / "变更登记一般情形审查规则表.xlsx",
    ROOT / "审查规则" / "变更登记一般情形审查规则表.xlsx",
]
SOURCE = next((path for path in SOURCE_CANDIDATES if path.exists()), SOURCE_CANDIDATES[0])
OUTPUT = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "rules"
    / "rules_change_general.json"
)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_dimension(value: str) -> str:
    raw = clean(value)
    if not raw:
        return "其他"
    raw = re.sub(r"^\d+\s*[.．、]\s*", "", raw).strip()
    mapping = {
        "格式模板": "格式模板规则库",
        "登记必填要素": "登记必填要素规则库",
        "跨材料数据逻辑": "跨材料数据逻辑校验库",
        "监管合规红线": "监管合规红线规则库",
        "业务退回/整改": "业务退回/整改案例库",
        "审查风险分级": "审查风险分级规则库",
    }
    return mapping.get(raw, raw)


def basis_file(basis_text: str) -> str:
    text = clean(basis_text)
    if not text:
        return ""
    head = re.split(r"[：:；;\n]", text, maxsplit=1)[0].strip()
    return head[:80]


def risk_level(rule_text: str, dimension: str, check_type: str) -> str:
    text = f"{rule_text} {dimension} {check_type}"
    if any(kw in text for kw in ("必须", "不得", "禁止", "必填", "缺失", "应完整", "应能有效佐证", "一致")):
        return "高风险"
    if "监管" in text or "跨材料" in text:
        return "中风险"
    return "中风险"


def file_check_type(rule_name: str, rule_text: str, dimension: str) -> str:
    text = f"{rule_name} {rule_text} {dimension}"
    if "佐证" in text or "匹配" in text or "一致" in text:
        return "跨材料一致性"
    if "材料" in text or "申请书" in text or "文件" in text:
        return "材料完整性"
    if "监管" in dimension:
        return "监管口径符合性"
    return "语义条件判断"


def element_check_type(raw: str, rule_text: str, dimension: str) -> str:
    value = clean(raw)
    if value:
        return value
    return file_check_type("", rule_text, dimension)


def file_ai_focus(rule_name: str, rule_text: str) -> list[str]:
    text = f"{rule_name} {rule_text}"
    if "佐证" in text or "证明" in text:
        return [
            "核对证明发生变更事实的文件是否能支撑本次变更事项",
            "结合变更登记申请书和本次变更登记申报模板判断变更事项是否匹配",
            "仅按本批给定规则判断，不做上一次登记模板与本次模板的差异审查",
        ]
    return [
        "核对变更登记申请书是否完整记载变更事项和变更原因",
        "指出申请书中缺失、不完整或与本次变更材料不匹配的位置",
        "仅按本批给定规则判断，不扩展审核范围",
    ]


def element_ai_focus(check_type: str) -> list[str]:
    if "跨材料" in check_type or "一致" in check_type:
        return [
            "以本次变更登记申报模板为字段填报数据源",
            "按规则明确要求核对申请书、证明文件与本次模板的一致性",
            "上一次登记申报模板仅作历史参考，除非规则明确要求，不得据此输出差异问题",
        ]
    return [
        "以本次变更登记申报模板为字段填报数据源",
        "只检查本批规则明确列出的字段、条件和监管口径",
        "上一次登记申报模板仅作历史参考，不得自行扩展为差异审查",
    ]


def file_applicable_materials(rule_text: str) -> list[str]:
    if "证明" in rule_text or "佐证" in rule_text:
        return ["申报模板", "申请书", "证明发生变更事实的文件", "上一次登记申报模板"]
    return ["申请书", "申报模板", "证明发生变更事实的文件"]


def element_applicable_materials(check_type: str, rule_text: str) -> list[str]:
    text = f"{check_type} {rule_text}"
    if "跨材料" in text or "一致" in text or "申请书" in text or "证明" in text or "信托文件" in text:
        return ["申报模板", "申请书", "证明发生变更事实的文件", "上一次登记申报模板"]
    return ["申报模板"]


def file_review_rules(ws) -> list[dict[str, Any]]:
    headers = [clean(c.value) for c in ws[1]]
    rules: list[dict[str, Any]] = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        data = {headers[i]: clean(row[i]) if i < len(row) else "" for i in range(len(headers))}
        if data.get("分类", "").lower() != "ai":
            continue
        rule_text = data.get("具体的规则", "")
        if not rule_text:
            continue
        idx = len(rules) + 1
        rule_name = data.get("文件审查规则", "") or data.get("文件名", "") or "变更登记文件审查"
        dim = normalize_dimension(data.get("审核维度", ""))
        ctype = file_check_type(rule_name, rule_text, dim)
        rules.append(
            {
                "rule_id": f"CHG-FILE-AI-{idx:03d}",
                "registration_type": "变更登记（一般情形）",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": row_no,
                "rule_name": rule_name,
                "rule_text": rule_text,
                "basis_file": basis_file(data.get("审核依据", "")),
                "basis_text": data.get("审核依据", ""),
                "review_dimension": dim,
                "table_name": "",
                "field_name": "",
                "applicable_materials": file_applicable_materials(rule_text),
                "review_method": "ai",
                "original_review_method": "AI",
                "check_type": ctype,
                "trigger_condition": data.get("文件名", ""),
                "machine_params": "rule_scope=given_rules_only; previous_template_reference_only=true",
                "risk_level": risk_level(rule_text, dim, ctype),
                "ai_check_focus": file_ai_focus(rule_name, rule_text),
                "evidence_requirement": "请指出涉及的材料、表名、字段、页码或文本片段。",
                "rule_role": "file_review",
                "severity_type": "issue",
                "issue_type": "文件审查",
                "field_anchor": "",
                "skip_when": "",
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def element_rules(ws) -> list[dict[str, Any]]:
    headers = [clean(c.value) for c in ws[1]]
    rules: list[dict[str, Any]] = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        data = {headers[i]: clean(row[i]) if i < len(row) else "" for i in range(len(headers))}
        method = data.get("执行方式", "")
        if method not in {"AI", "ai", "混合"}:
            continue
        rule_text = data.get("具体规则", "")
        if not rule_text:
            continue
        idx = len(rules) + 1
        table_name = data.get("表名", "")
        field_name = data.get("要素名", "")
        dim = normalize_dimension(data.get("审查维度", ""))
        ctype = element_check_type(data.get("校验类型", ""), rule_text, dim)
        rules.append(
            {
                "rule_id": f"CHG-ELEMENT-AI-{idx:03d}",
                "registration_type": "变更登记（一般情形）",
                "rule_source": "内置规则",
                "source_file": SOURCE.name,
                "source_sheet": ws.title,
                "source_row": row_no,
                "rule_name": field_name or table_name or "变更登记要素审查",
                "rule_text": rule_text,
                "basis_file": basis_file(data.get("审核依据", "")),
                "basis_text": data.get("审核依据", ""),
                "review_dimension": dim,
                "table_name": table_name,
                "field_name": field_name,
                "applicable_materials": element_applicable_materials(ctype, rule_text),
                "review_method": "ai",
                "original_review_method": method,
                "check_type": ctype,
                "trigger_condition": data.get("触发条件", ""),
                "machine_params": data.get("机器参数", ""),
                "risk_level": risk_level(rule_text, dim, ctype),
                "ai_check_focus": element_ai_focus(ctype),
                "evidence_requirement": "请指出涉及的材料、表名、字段、页码或文本片段。",
                "rule_role": "field_primary",
                "severity_type": "issue",
                "issue_type": "要素审查",
                "field_anchor": field_name,
                "skip_when": "",
                "enabled": True,
                "demo_enabled": True,
            }
        )
    return rules


def main() -> None:
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    rules = file_review_rules(wb["文件审查"]) + element_rules(wb["要素审查"])
    payload = {
        "rule_set_id": "change_general_ai_rules_v1",
        "registration_type": "变更登记（一般情形）",
        "source_file": SOURCE.name,
        "description": "变更登记（一般情形）AI 审核规则；来源于《变更登记一般情形审查规则表.xlsx》的文件审查 ai 规则和要素审查 AI/混合规则。",
        "total_rules": len(rules),
        "rules": rules,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {OUTPUT} ({len(rules)} rules)")


if __name__ == "__main__":
    main()
