from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

import fitz
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from backend.app import app
from backend.services import material_store
from backend.services.material_parser import parse_material


class MaterialParserTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_json_txt_docx_and_pdf_parsing(self):
        json_path = self.root / "预登记申报模板.json"
        json_path.write_text(
            json.dumps({"产品基本信息": {"产品名称": "测试计划"}, "记录": [{"金额": 10}, {"金额": 20}]}, ensure_ascii=False),
            encoding="utf-8",
        )
        parsed_json = parse_material(json_path)
        self.assertEqual(parsed_json.file_kind, "json")
        self.assertEqual(
            {(item.location, item.text) for item in parsed_json.segments},
            {
                ("产品基本信息.产品名称", "测试计划"),
                ("记录[0].金额", "10"),
                ("记录[1].金额", "20"),
            },
        )

        txt_path = self.root / "申请书.txt"
        txt_path.write_text("第一段\n\n第二段", encoding="utf-8")
        parsed_txt = parse_material(txt_path)
        self.assertEqual([item.text for item in parsed_txt.segments], ["第一段", "第二段"])

        docx_path = self.root / "信托合同.docx"
        document = Document()
        document.add_paragraph("合同正文")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "内容"
        document.save(docx_path)
        parsed_docx = parse_material(docx_path)
        self.assertEqual(parsed_docx.file_kind, "docx")
        self.assertTrue(any(item.text == "合同正文" for item in parsed_docx.segments))
        self.assertTrue(any("字段 | 内容" in item.text for item in parsed_docx.segments))

        pdf_path = self.root / "申请书.pdf"
        pdf = fitz.open()
        text_page = pdf.new_page()
        text_page.insert_text((72, 72), "Trust review material")
        pdf.new_page()
        pdf.save(pdf_path)
        pdf.close()
        parsed_pdf = parse_material(pdf_path)
        self.assertEqual(len(parsed_pdf.segments), 2)
        self.assertIn("Trust review material", parsed_pdf.segments[0].text)
        self.assertIn("不支持OCR", parsed_pdf.segments[1].text)

    def test_metadata_and_generic_excel_parsing(self):
        template_path = self.root / "初始登记申报模板.xlsx"
        workbook = Workbook()
        metadata = workbook.active
        metadata.title = "要素表"
        product = workbook.create_sheet("产品要素")
        grid = workbook.create_sheet("共同受托人信息")
        metadata.cell(2, 2, "1.产品基本信息")
        metadata.cell(3, 2, "产品名称")
        metadata.cell(3, 6, "B")
        metadata.cell(3, 7, 2)
        metadata.cell(4, 2, "2.共同受托人信息")
        metadata.cell(5, 2, "共同受托人名称")
        metadata.cell(5, 6, "A")
        metadata.cell(5, 7, 5)
        metadata.cell(5, 14, "共同受托人信息")
        product["B2"] = "测试信托计划"
        grid["A5"] = "甲公司"
        grid["A6"] = "乙公司"
        workbook.save(template_path)
        workbook.close()

        parsed = parse_material(template_path)
        values = {(item.location, item.text) for item in parsed.segments}
        self.assertIn(("产品基本信息.产品名称", "测试信托计划"), values)
        self.assertIn(("共同受托人信息[0].共同受托人名称", "甲公司"), values)
        self.assertIn(("共同受托人信息[1].共同受托人名称", "乙公司"), values)

        generic_path = self.root / "普通数据.xlsx"
        workbook = Workbook()
        workbook.active.append(["名称", "金额"])
        workbook.active.append(["测试", 10])
        workbook.save(generic_path)
        workbook.close()
        generic = parse_material(generic_path)
        self.assertEqual(generic.segments[1].text, "测试 | 10")

    def test_unsupported_format_rejected(self):
        path = self.root / "材料.csv"
        path.write_text("a,b", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "暂不支持"):
            parse_material(path)


class MaterialApiTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_upload_dir = material_store.MATERIAL_UPLOAD_DIR
        material_store.MATERIAL_UPLOAD_DIR = Path(self.temp_dir.name) / "materials"
        self.client = TestClient(app)

    def tearDown(self):
        material_store.MATERIAL_UPLOAD_DIR = self.original_upload_dir
        self.temp_dir.cleanup()

    def upload_json(self, process: str = "pre_registration", filename: str = "预登记申报模板.json"):
        return self.client.post(
            "/api/rule-engine/materials",
            data={"process": process},
            files={"file": (filename, b'{"product":{"name":"demo"}}', "application/json")},
        )

    def test_upload_list_detail_extracted_duplicate_and_delete(self):
        uploaded = self.upload_json()
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        metadata = uploaded.json()
        self.assertEqual(metadata["parse_status"], "已解析")
        self.assertEqual(metadata["segments_count"], 1)
        file_id = metadata["file_id"]

        listed = self.client.get("/api/rule-engine/materials", params={"process": "pre_registration"})
        self.assertEqual([item["file_id"] for item in listed.json()], [file_id])
        self.assertEqual(self.client.get(f"/api/rule-engine/materials/{file_id}").status_code, 200)
        extracted = self.client.get(f"/api/rule-engine/materials/{file_id}/extracted")
        self.assertEqual(extracted.status_code, 200)
        self.assertEqual(extracted.json()["material_name"], "预登记申报模板.json")
        self.assertEqual(extracted.json()["size_bytes"], len(b'{"product":{"name":"demo"}}'))

        duplicate = self.upload_json()
        self.assertEqual(duplicate.status_code, 409)
        other_process = self.upload_json(process="initial")
        self.assertEqual(other_process.status_code, 200)

        deleted = self.client.delete(f"/api/rule-engine/materials/{file_id}")
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(self.client.get(f"/api/rule-engine/materials/{file_id}").status_code, 404)

    def test_invalid_format_parse_failure_and_size_limit(self):
        unsupported = self.client.post(
            "/api/rule-engine/materials",
            data={"process": "pre_registration"},
            files={"file": ("材料.csv", b"a,b", "text/csv")},
        )
        self.assertEqual(unsupported.status_code, 400)

        invalid_json = self.client.post(
            "/api/rule-engine/materials",
            data={"process": "pre_registration"},
            files={"file": ("损坏模板.json", b"not-json", "application/json")},
        )
        self.assertEqual(invalid_json.status_code, 200)
        self.assertEqual(invalid_json.json()["parse_status"], "解析失败")
        extracted = self.client.get(
            f"/api/rule-engine/materials/{invalid_json.json()['file_id']}/extracted"
        )
        self.assertEqual(extracted.status_code, 409)

        from backend.api import material_router

        original_limit = material_router.MAX_MATERIAL_UPLOAD_BYTES
        material_router.MAX_MATERIAL_UPLOAD_BYTES = 5
        try:
            too_large = self.client.post(
                "/api/rule-engine/materials",
                data={"process": "pre_registration"},
                files={"file": ("超限模板.json", b"123456", "application/json")},
            )
            self.assertEqual(too_large.status_code, 413)
        finally:
            material_router.MAX_MATERIAL_UPLOAD_BYTES = original_limit


if __name__ == "__main__":
    unittest.main()
