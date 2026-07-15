from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import app
from backend.services import rule_library


class RuleLibraryApiTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_file = rule_library.RULE_LIBRARY_FILE
        rule_library.RULE_LIBRARY_FILE = Path(self.temp_dir.name) / "rules.json"
        self.client = TestClient(app)

    def tearDown(self):
        rule_library.RULE_LIBRARY_FILE = self.original_file
        self.temp_dir.cleanup()

    def script_payload(self, rule_id: str = "INITIAL-SCRIPT-001"):
        return {
            "rule_id": rule_id,
            "rule_name": "募集规模必须填写",
            "process": "initial",
            "review_method": "script",
            "applicable_materials": ["初始登记申报模板"],
            "rule_object": "字段",
            "table_name": "产品基本信息",
            "field_path": "产品基本信息.募集规模",
            "review_dimension": "法定要素完整性",
            "trigger_condition": "",
            "rule_text": "募集规模为必填项。",
            "basis_text": "初始登记要素表及填表说明。",
            "risk_level": "中风险",
            "version": "V1.0",
            "enabled": True,
            "remarks": "",
            "operator": "required",
            "script_params": {},
            "special_prompt": "",
        }

    def test_empty_library_isolated_by_process(self):
        for process in ("pre_registration", "pre_report", "initial", "termination"):
            response = self.client.get("/api/rule-library/rules", params={"process": process})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["rules"], [])

    def test_script_rule_crud_copy_and_toggle(self):
        created = self.client.post("/api/rule-library/rules", json=self.script_payload())
        self.assertEqual(created.status_code, 200, created.text)
        self.assertIn("不得为空", created.json()["visual_rule"])

        duplicate = self.client.post("/api/rule-library/rules", json=self.script_payload())
        self.assertEqual(duplicate.status_code, 409)

        copied = self.client.post("/api/rule-library/rules/INITIAL-SCRIPT-001/copy")
        self.assertEqual(copied.status_code, 200, copied.text)
        self.assertFalse(copied.json()["enabled"])
        self.assertNotEqual(copied.json()["rule_id"], "INITIAL-SCRIPT-001")

        toggled = self.client.patch(
            "/api/rule-library/rules/INITIAL-SCRIPT-001/enabled", json={"enabled": False}
        )
        self.assertEqual(toggled.status_code, 200)
        self.assertFalse(toggled.json()["enabled"])

        listed = self.client.get("/api/rule-library/rules", params={"process": "initial"})
        self.assertEqual(listed.json()["total"], 2)

        copied_id = copied.json()["rule_id"]
        deleted = self.client.post(
            "/api/rule-library/rules/batch-delete", json={"rule_ids": [copied_id]}
        )
        self.assertEqual(deleted.status_code, 200, deleted.text)
        self.assertEqual(deleted.json()["deleted_count"], 1)

        removed = self.client.delete("/api/rule-library/rules/INITIAL-SCRIPT-001")
        self.assertEqual(removed.status_code, 200)

    def _ai_workbook(self, rule_id: str = "PRE-AI-001") -> bytes:
        content = rule_library.build_import_template("pre_report")
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook["AI审核规则"]
        headers = {cell.value: index for index, cell in enumerate(sheet[1], start=1)}
        values = {
            "规则ID": rule_id,
            "规则名称": "禁止保本保收益表述",
            "适用流程": "事前报告",
            "审核方式": "AI规则",
            "适用材料": "事前报告申请书、信托合同",
            "规则对象": "跨材料",
            "要素分类/表名": "",
            "字段路径": "",
            "审查维度": "文本语义合规",
            "触发条件": "",
            "具体规则": "材料不得承诺保本保收益。",
            "审查依据": "信托登记审查规则。",
            "风险等级": "高风险",
            "当前版本": "V1.0",
            "启用状态": "启用",
            "备注": "",
            "专属提示词": "识别明示或暗示的保本保收益表述，必须引用原文证据。",
        }
        for header, value in values.items():
            sheet.cell(2, headers[header], value)
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()

    def test_excel_preview_commit_and_duplicate_rejection(self):
        content = self._ai_workbook()
        files = {
            "file": (
                "事前报告规则.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }
        preview = self.client.post(
            "/api/rule-library/import/preview", params={"process": "pre_report"}, files=files
        )
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertTrue(preview.json()["valid"])
        self.assertEqual(preview.json()["ai_count"], 1)

        committed = self.client.post(
            "/api/rule-library/import/commit", params={"process": "pre_report"}, files=files
        )
        self.assertEqual(committed.status_code, 200, committed.text)
        self.assertEqual(committed.json()["imported_count"], 1)

        duplicate_preview = self.client.post(
            "/api/rule-library/import/preview", params={"process": "pre_report"}, files=files
        )
        self.assertEqual(duplicate_preview.status_code, 200)
        self.assertFalse(duplicate_preview.json()["valid"])
        self.assertIn("已存在于规则库", duplicate_preview.json()["errors"][0]["message"])

        duplicate_commit = self.client.post(
            "/api/rule-library/import/commit", params={"process": "pre_report"}, files=files
        )
        self.assertEqual(duplicate_commit.status_code, 400)
        listed = self.client.get("/api/rule-library/rules", params={"process": "pre_report"})
        self.assertEqual(listed.json()["total"], 1)

    def test_template_contains_descriptions_and_validation_sheets(self):
        workbook = load_workbook(io.BytesIO(rule_library.build_import_template("initial")))
        self.assertEqual(
            workbook.sheetnames,
            ["脚本审核规则", "AI审核规则", "列说明", "operator字典"],
        )
        self.assertEqual(workbook["列说明"]["A2"].value, "规则ID")
        self.assertIsNotNone(workbook["脚本审核规则"]["A1"].comment)
        self.assertGreater(len(workbook["脚本审核规则"].data_validations.dataValidation), 0)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
