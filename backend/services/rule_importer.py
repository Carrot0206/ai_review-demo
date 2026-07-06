"""规则表格导入：把 Excel 行转换为用户新增规则。"""
from __future__ import annotations

import io
import time
import uuid
from typing import Any

from openpyxl import load_workbook

from .schemas import PROCESS_LABEL, ProcessType, RiskLevel
from .user_rules import UserRule


HEADER_ALIASES = {
    "process": ["流程", "登记流程", "适用流程"],
    "material": ["文件名", "文件/子项", "适用材料", "材料类型", "材料名称"],
    "table_name": ["表名", "要素分类", "表单", "模块"],
    "field_name": ["要素名", "字段", "字段名", "数据项", "英文名"],
    "review_dimension": ["审查维度", "审核维度", "规则类别", "分类结果"],
    "rule_text": [
        "具体规则",
        "具体的规则",
        "具体规则(已融合全部枚举值)",
        "文件审查规则",
        "审查规则",
        "规则内容",
    ],
    "basis_text": [
        "审核依据",
        "审查依据",
        "审核依据（法规或者文件里规定怎么写）",
        "字段依据",
    ],
    "check_type": ["校验类型", "执行方式", "分类", "检查类型"],
    "trigger_condition": ["触发条件", "适用条件"],
    "risk_level": ["风险等级", "风险级别", "风险"],
}

CHECK_TYPE_MAP = {
    "需AI推理判断": "语义条件判断",
    "AI推理判断": "语义条件判断",
    "文件审查": "语义条件判断",
    "要素审查": "语义条件判断",
    "可直接比对": "字段类型格式",
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _norm_header(value: Any) -> str:
    return _text(value).replace("\n", "").replace(" ", "")


def _find_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]]:
    best_idx = -1
    best_map: dict[str, int] = {}
    best_score = 0
    alias_lookup = {
        _norm_header(alias): key
        for key, aliases in HEADER_ALIASES.items()
        for alias in aliases
    }
    for idx, row in enumerate(rows[:10]):
        mapping: dict[str, int] = {}
        for col, cell in enumerate(row):
            key = alias_lookup.get(_norm_header(cell))
            if key and key not in mapping:
                mapping[key] = col
        score = len(mapping)
        if "rule_text" in mapping:
            score += 4
        if "review_dimension" in mapping:
            score += 1
        if score > best_score:
            best_idx = idx
            best_map = mapping
            best_score = score
    if "rule_text" not in best_map:
        return -1, {}
    return best_idx, best_map


def _cell(row: tuple[Any, ...], mapping: dict[str, int], key: str) -> str:
    col = mapping.get(key)
    if col is None or col >= len(row):
        return ""
    return _text(row[col])


def _split_materials(value: str) -> list[str]:
    if not value:
        return []
    import re

    parts = [p.strip() for p in re.split(r"[、,，;；/\n]+", value) if p.strip()]
    return parts[:8]


def _risk(value: str) -> RiskLevel:
    if "高" in value:
        return "高风险"
    if "低" in value:
        return "低风险"
    return "中风险"


def _check_type(value: str) -> str:
    if not value:
        return "语义条件判断"
    return CHECK_TYPE_MAP.get(value, value)


def _row_process_matches(row_process: str, process: ProcessType) -> bool:
    if not row_process:
        return True
    target = PROCESS_LABEL.get(process, process)
    return row_process == target or target in row_process or row_process in target


def import_rules_from_excel(
    content: bytes,
    *,
    filename: str,
    process: ProcessType,
) -> list[UserRule]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    now = int(time.time())
    imported: list[UserRule] = []
    seen: set[tuple[str, str, str, str]] = set()

    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            header_idx, mapping = _find_header_row(rows)
            if header_idx < 0:
                continue

            for source_row, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
                rule_text = _cell(row, mapping, "rule_text")
                if not rule_text:
                    continue
                row_process = _cell(row, mapping, "process")
                if not _row_process_matches(row_process, process):
                    continue

                table_name = _cell(row, mapping, "table_name")
                field_name = _cell(row, mapping, "field_name")
                material = _cell(row, mapping, "material")
                dimension = _cell(row, mapping, "review_dimension") or "用户新增规则"
                basis = _cell(row, mapping, "basis_text")
                if basis and basis not in rule_text:
                    rule_text = f"{rule_text}\n依据：{basis}"

                fingerprint = (ws.title, table_name, field_name, rule_text)
                if fingerprint in seen:
                    continue
                seen.add(fingerprint)

                imported.append(
                    UserRule(
                        rule_id=f"USER-{uuid.uuid4().hex[:8].upper()}",
                        process=process,
                        applicable_materials=_split_materials(material),
                        rule_text=rule_text,
                        risk_level=_risk(_cell(row, mapping, "risk_level")),
                        review_dimension=dimension,
                        check_type=_check_type(_cell(row, mapping, "check_type")),
                        table_name=table_name,
                        field_name=field_name,
                        trigger_condition=_cell(row, mapping, "trigger_condition"),
                        created_at=now,
                        enabled=True,
                    )
                )
    finally:
        wb.close()

    return imported
