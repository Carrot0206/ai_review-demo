"""规则版本 Excel 导入。

读取导入版 Excel 的「脚本审核规则」「AI审核规则」sheet，
转换为现有 review_service 使用的 Rule 结构。
两个规则 sheet 至少提供一个；允许只上传 AI 审核规则。
"""
from __future__ import annotations

import io
import json
from typing import Any

from openpyxl import load_workbook

from .schemas import PROCESS_LABEL, ProcessType, Rule, TableScopeRule


SCRIPT_SHEET = "脚本审核规则"
AI_SHEET = "AI审核规则"
SCOPE_SHEET = "表级报送范围"


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _headers(ws) -> dict[str, int]:
    return {_text(cell.value): idx for idx, cell in enumerate(ws[1]) if _text(cell.value)}


def _cell(row: tuple[Any, ...], mapping: dict[str, int], key: str) -> str:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return ""
    return _text(row[idx])


def _split(value: str) -> list[str]:
    if not value:
        return []
    import re

    return [p.strip() for p in re.split(r"[、,，;；/\n]+", value) if p.strip()]


def _risk(value: str) -> str:
    if "高" in value:
        return "高风险"
    if "低" in value:
        return "低风险"
    return "中风险"


def _normalize_table_name(value: str) -> str:
    import re

    text = re.sub(r"^\s*\d+\s*[.、．]\s*", "", value or "").strip()
    if text in {"关联交易事项", "关联交易信息"}:
        return "关联交易信息"
    return text


def _parse_params(raw: str, errors: list[dict], *, sheet: str, row_no: int, rule_id: str) -> dict:
    if not raw:
        return {}
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as e:
        errors.append(
            {
                "sheet": sheet,
                "row": row_no,
                "rule_id": rule_id,
                "message": f"参数(JSON)解析失败：{e}",
            }
        )
        return {}
    if not isinstance(parsed, dict):
        errors.append(
            {
                "sheet": sheet,
                "row": row_no,
                "rule_id": rule_id,
                "message": "参数(JSON)必须是对象",
            }
        )
        return {}
    return parsed


def _normalize_dimension(value: str) -> str:
    mapping = {
        "格式模板": "格式模板规则库",
        "登记必填要素": "登记必填要素规则库",
        "跨材料数据逻辑": "跨材料数据逻辑校验库",
        "监管合规红线": "监管合规红线规则库",
        "业务退回/整改": "业务退回/整改案例库",
        "审查风险分级": "审查风险分级规则库",
    }
    return mapping.get(value, value or "其他")


def infer_rule_type(rule: Rule) -> str:
    """Infer a conservative prompt profile for older rule workbooks."""
    if rule.rule_type:
        return rule.rule_type
    text = " ".join(
        str(value or "")
        for value in [
            rule.check_type,
            rule.operator,
            rule.rule_text,
            rule.rule_name,
            rule.review_dimension,
            " ".join(rule.ai_check_focus or []),
        ]
    )
    if "版式" in text or "显著位置" in text:
        return "human_layout"
    if rule.rule_role == "script_ai_fallback":
        return "script_ai_fallback"
    if rule.review_method == "script":
        return ""
    if "跨材料一致性" in text or "保持一致" in text or "与信托文件" in text or "与申请书" in text:
        return "cross_material_consistency"
    if "主键" in text or "唯一性" in text or "唯一约束" in text:
        return "array_unique"
    if any(keyword in text for keyword in ("规模", "金额", "上限", "下限", "比例", "加总", "应为0", "不小于", "不大于")):
        return "numeric_relation"
    if any(keyword in text for keyword in ("分类", "监管口径", "字段联动", "系统提示分类错误", "不得分类为")):
        return "business_mapping"
    if any(keyword in text for keyword in ("填表范围", "非必填", "仅要求", "仅当", "选填")):
        return "form_scope"
    if any(keyword in text for keyword in ("须提交", "应提交", "上传", "材料", "清晰", "可阅读", "佐证")):
        return "material_required"
    return "general_explanation"


def _script_rule_needs_configuration(rule: Rule) -> bool:
    if rule.review_method != "script":
        return False
    if rule.operator in {"custom_condition", "format_check"}:
        return True
    if rule.operator != "conditional_compare":
        return False
    params = rule.script_params or {}
    if params.get("external_source_required"):
        return True
    target = params.get("target")
    return not isinstance(target, dict) or not target.get("field") or not target.get("op")


def script_rule_needs_configuration(rule: Rule) -> bool:
    """Public helper used by review flow for AI fallback."""
    return _script_rule_needs_configuration(rule)


def _script_rule_configuration_error(rule: Rule) -> str:
    if rule.review_method != "script" or rule.operator != "dependent_enum":
        return ""
    params = rule.script_params or {}
    if not str(params.get("parent_field") or "").strip():
        return "dependent_enum 缺少 parent_field"
    if not str(params.get("child_field") or "").strip():
        return "dependent_enum 缺少 child_field"
    mapping = params.get("mapping")
    if not isinstance(mapping, dict) or not mapping:
        return "dependent_enum 缺少非空 mapping 对象"
    if any(not isinstance(values, list) or not values for values in mapping.values()):
        return "dependent_enum 的 mapping 每一项都必须是非空数组"
    return ""


def _script_rule_configuration_warning(rule: Rule) -> str:
    if rule.review_method != "script" or rule.operator != "enum":
        return ""
    values = (rule.script_params or {}).get("values")
    if not isinstance(values, list) or not any(str(value or "").strip() for value in values):
        return "enum 缺少非空 values，规则将跳过且不会转入 AI 兜底"
    return ""


def _rule_common(
    *,
    process: ProcessType,
    filename: str,
    row: tuple[Any, ...],
    mapping: dict[str, int],
    sheet: str,
    row_no: int,
    review_method: str,
    errors: list[dict],
) -> Rule | None:
    rule_id = _cell(row, mapping, "规则ID")
    rule_text = _cell(row, mapping, "具体规则")
    if not rule_id:
        errors.append({"sheet": sheet, "row": row_no, "message": "缺少规则ID"})
        return None
    if not rule_text:
        errors.append({"sheet": sheet, "row": row_no, "rule_id": rule_id, "message": "缺少具体规则"})
        return None

    field_path = _cell(row, mapping, "字段路径")
    section = _cell(row, mapping, "要素分类/表名")
    field_name = _cell(row, mapping, "规则名称/要素名")
    if "." in field_path:
        section, field_name = field_path.rsplit(".", 1)
    elif field_path:
        field_name = field_path

    source_sheet = _cell(row, mapping, "来源sheet") or sheet
    source_row_raw = _cell(row, mapping, "来源行")
    try:
        source_row = int(source_row_raw) if source_row_raw else row_no
    except ValueError:
        source_row = row_no

    operator = _cell(row, mapping, "operator")
    params = _parse_params(
        _cell(row, mapping, "参数(JSON)"),
        errors,
        sheet=sheet,
        row_no=row_no,
        rule_id=rule_id,
    )
    if review_method == "script" and not operator:
        errors.append({"sheet": sheet, "row": row_no, "rule_id": rule_id, "message": "脚本规则缺少operator"})

    return Rule(
        rule_id=rule_id,
        registration_type=PROCESS_LABEL.get(process, process),
        rule_source="上传规则",
        source_file=filename,
        source_sheet=source_sheet,
        source_row=source_row,
        rule_name=_cell(row, mapping, "规则名称/要素名") or rule_id,
        rule_text=rule_text,
        basis_file=_cell(row, mapping, "审查依据") or filename,
        basis_text=_cell(row, mapping, "审查依据"),
        review_dimension=_normalize_dimension(_cell(row, mapping, "审查维度")),
        rule_type=_cell(row, mapping, "规则类型"),
        table_name=section,
        field_name=field_name,
        applicable_materials=_split(_cell(row, mapping, "适用材料")),
        review_method=review_method,
        rule_object=_cell(row, mapping, "规则对象"),
        operator=operator,
        script_params=params,
        check_type=_cell(row, mapping, "比对方式") or _cell(row, mapping, "AI审核重点") or operator or "语义条件判断",
        trigger_condition=_cell(row, mapping, "触发条件"),
        machine_params=json.dumps({"operator": operator, "params": params}, ensure_ascii=False) if operator else "",
        risk_level=_risk(_cell(row, mapping, "风险等级")),
        ai_check_focus=[_cell(row, mapping, "AI审核重点")] if _cell(row, mapping, "AI审核重点") else [],
        evidence_requirement="请指出涉及的材料、表名、字段、页码或文本片段。",
        rule_role="uploaded_script" if review_method == "script" else "uploaded_ai",
        severity_type="issue",
        issue_type=_cell(row, mapping, "规则对象") or "",
        field_anchor=field_path or field_name,
        enabled=_cell(row, mapping, "启用状态") not in {"否", "false", "False", "0"},
        demo_enabled=True,
    )
    if not rule.rule_type:
        rule.rule_type = infer_rule_type(rule)
    return rule


def _scope_rules(
    wb,
    *,
    filename: str,
    process: ProcessType,
    errors: list[dict],
    warnings: list[dict],
) -> list[TableScopeRule]:
    if SCOPE_SHEET not in wb.sheetnames:
        warnings.append(
            {
                "sheet": SCOPE_SHEET,
                "row": 0,
                "message": "未配置表级报送范围，按旧版规则执行",
            }
        )
        return []

    ws = wb[SCOPE_SHEET]
    mapping = _headers(ws)
    for required in ("范围规则ID", "表名", "范围类型", "具体规则"):
        if required not in mapping:
            raise ValueError(f"{SCOPE_SHEET} 缺少必要列：{required}")

    scopes: list[TableScopeRule] = []
    valid_scope_types = {"always_required", "conditional", "prohibited"}
    valid_outside = {"not_required", "prohibited", "unknown"}
    valid_unknown = {"ai_review", "human_review"}
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(_text(value) for value in row):
            continue
        scope_id = _cell(row, mapping, "范围规则ID")
        table_name = _normalize_table_name(_cell(row, mapping, "表名"))
        scope_type = _cell(row, mapping, "范围类型") or "conditional"
        rule_text = _cell(row, mapping, "具体规则")
        if not scope_id or not table_name or not rule_text:
            errors.append(
                {
                    "sheet": SCOPE_SHEET,
                    "row": row_no,
                    "rule_id": scope_id,
                    "message": "范围规则ID、表名和具体规则均为必填项",
                }
            )
            continue
        if scope_type not in valid_scope_types:
            errors.append(
                {
                    "sheet": SCOPE_SHEET,
                    "row": row_no,
                    "rule_id": scope_id,
                    "message": f"不支持的范围类型：{scope_type}",
                }
            )
            continue
        applicability = _parse_params(
            _cell(row, mapping, "适用条件(JSON)"),
            errors,
            sheet=SCOPE_SHEET,
            row_no=row_no,
            rule_id=scope_id,
        )
        exemption = _parse_params(
            _cell(row, mapping, "豁免条件(JSON)"),
            errors,
            sheet=SCOPE_SHEET,
            row_no=row_no,
            rule_id=scope_id,
        )
        outside_effect = _cell(row, mapping, "区间外处理") or "unknown"
        unknown_policy = _cell(row, mapping, "未知处理") or "ai_review"
        if outside_effect not in valid_outside:
            errors.append({"sheet": SCOPE_SHEET, "row": row_no, "rule_id": scope_id, "message": f"不支持的区间外处理：{outside_effect}"})
            continue
        if unknown_policy not in valid_unknown:
            errors.append({"sheet": SCOPE_SHEET, "row": row_no, "rule_id": scope_id, "message": f"不支持的未知处理：{unknown_policy}"})
            continue
        scopes.append(
            TableScopeRule(
                scope_id=scope_id,
                registration_type=PROCESS_LABEL.get(process, process),
                source_file=filename,
                source_row=row_no,
                table_name=table_name,
                scope_type=scope_type,
                applicability_condition=applicability,
                exemption_condition=exemption,
                ai_materials=_split(_cell(row, mapping, "AI判定材料")),
                ai_check_focus=_cell(row, mapping, "AI判定重点"),
                effective_from=_cell(row, mapping, "生效日期"),
                effective_to=_cell(row, mapping, "失效日期"),
                outside_effect=outside_effect,
                unknown_policy=unknown_policy,
                rule_text=rule_text,
                basis_text=_cell(row, mapping, "审查依据"),
                risk_level=_risk(_cell(row, mapping, "风险等级")),
                enabled=_cell(row, mapping, "启用状态") not in {"否", "false", "False", "0"},
            )
        )
    return scopes


def import_rule_package_from_excel(
    content: bytes,
    *,
    filename: str,
    process: ProcessType,
) -> tuple[list[Rule], list[TableScopeRule], dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rules: list[Rule] = []
    errors: list[dict] = []
    warnings: list[dict] = []
    try:
        available_rule_sheets = [
            item
            for item in ((SCRIPT_SHEET, "script"), (AI_SHEET, "ai"))
            if item[0] in wb.sheetnames
        ]
        if not available_rule_sheets:
            raise ValueError(f"缺少必要sheet：{SCRIPT_SHEET} 或 {AI_SHEET}")

        for sheet, method in available_rule_sheets:
            ws = wb[sheet]
            mapping = _headers(ws)
            for required in ("规则ID", "具体规则"):
                if required not in mapping:
                    raise ValueError(f"{sheet} 缺少必要列：{required}")
            if method == "script" and "operator" not in mapping:
                raise ValueError(f"{sheet} 缺少必要列：operator")

            for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(_text(v) for v in row):
                    continue
                rule = _rule_common(
                    process=process,
                    filename=filename,
                    row=row,
                    mapping=mapping,
                    sheet=sheet,
                    row_no=row_no,
                    review_method=method,
                    errors=errors,
                )
                if rule is None:
                    continue
                configuration_error = _script_rule_configuration_error(rule)
                if configuration_error:
                    errors.append(
                        {
                            "sheet": sheet,
                            "row": row_no,
                            "rule_id": rule.rule_id,
                            "message": configuration_error,
                        }
                    )
                configuration_warning = _script_rule_configuration_warning(rule)
                if configuration_warning:
                    warnings.append(
                        {
                            "sheet": sheet,
                            "row": row_no,
                            "rule_id": rule.rule_id,
                            "message": configuration_warning,
                        }
                    )
                if _script_rule_needs_configuration(rule):
                    warnings.append(
                        {
                            "sheet": sheet,
                            "row": row_no,
                            "rule_id": rule.rule_id,
                            "message": f"{rule.operator} 暂按保守策略入库，参数不足时审核阶段跳过",
                        }
                    )
                rules.append(rule)
        scopes = _scope_rules(
            wb,
            filename=filename,
            process=process,
            errors=errors,
            warnings=warnings,
        )
    finally:
        wb.close()

    seen: set[str] = set()
    duplicates = []
    for r in rules:
        if r.rule_id in seen:
            duplicates.append(r.rule_id)
        seen.add(r.rule_id)
    if duplicates:
        raise ValueError(f"规则ID重复：{', '.join(sorted(set(duplicates)))}")
    scope_ids = [scope.scope_id for scope in scopes]
    if len(scope_ids) != len(set(scope_ids)):
        raise ValueError("表级报送范围存在重复的范围规则ID")

    report = {
        "total_rules": len(rules),
        "script_count": sum(1 for r in rules if r.review_method == "script"),
        "ai_count": sum(1 for r in rules if r.review_method == "ai"),
        "unsupported_count": sum(1 for r in rules if _script_rule_needs_configuration(r)),
        "error_count": len(errors),
        "warning_count": len(warnings),
        "scope_count": len(scopes),
        "errors": errors,
        "warnings": warnings,
    }
    if errors:
        first = errors[0]
        raise ValueError(
            f"{len(errors)} 条规则解析错误，首条：{first.get('sheet', '')} 第{first.get('row', '')}行 {first.get('message', '')}"
        )
    return rules, scopes, report


def import_rule_set_from_excel(content: bytes, *, filename: str, process: ProcessType) -> tuple[list[Rule], dict]:
    """Backward-compatible rule-only import used by existing validation tools."""
    rules, _scopes, report = import_rule_package_from_excel(
        content,
        filename=filename,
        process=process,
    )
    return rules, report
