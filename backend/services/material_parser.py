"""将上传材料解析为规则引擎统一的 ExtractedMaterial。"""
from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from ..models.schemas import ExtractedMaterial, MaterialSegment


SUPPORTED_EXTENSIONS = {".json", ".xlsx", ".xlsm", ".pdf", ".docx", ".txt"}

MATERIAL_TYPE_RULES = [
    ("上一次", "上一次登记申报模板"),
    ("上一笔", "上一次登记申报模板"),
    ("历史登记", "上一次登记申报模板"),
    ("证明发生需要更正事实", "证明发生需要更正事实的文件"),
    ("更正事实证明", "证明发生需要更正事实的文件"),
    ("证明发生变更事实", "证明发生变更事实的文件"),
    ("变更事实证明", "证明发生变更事实的文件"),
    ("证明文件", "证明发生变更事实的文件"),
    ("原预登记", "原预登记申报模板JSON"),
    ("baseline", "原预登记申报模板JSON"),
    ("Baseline", "原预登记申报模板JSON"),
    ("系统记录", "原预登记系统记录"),
    ("政信", "政信类证明材料"),
    ("融资平台债务", "政信类证明材料"),
    ("补充预登记申报模板", "申报模板"),
    ("补充预登记模板", "申报模板"),
    ("补充预登记申请书", "申请书"),
    ("补充说明", "补充说明材料"),
    ("要素报告表", "信托预登记要素报告表"),
    ("新型资产服务信托", "新型资产服务信托情况说明"),
    ("情况说明", "新型资产服务信托情况说明"),
    ("模板", "申报模板"),
    ("申请书", "申请书"),
    ("信托文件", "信托文件样本"),
    ("信托合同", "信托文件样本"),
    ("清算报告", "其他附件"),
]


def guess_material_type(filename: str) -> str:
    name = Path(filename).name
    for keyword, material_type in MATERIAL_TYPE_RULES:
        if keyword in name:
            return material_type
    if Path(name).suffix.lower() in {".xlsx", ".xlsm"}:
        return "申报模板"
    return "其他附件"


def _flatten_json(value: Any, prefix: str = "") -> list[tuple[str, str]]:
    output: list[tuple[str, str]] = []
    if isinstance(value, dict):
        for key, child in value.items():
            location = f"{prefix}.{key}" if prefix else str(key)
            output.extend(_flatten_json(child, location))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            output.extend(_flatten_json(child, f"{prefix}[{index}]"))
    else:
        output.append((prefix or "(root)", "" if value is None else str(value)))
    return output


def parse_json_template(path: Path, material_type: Optional[str] = None) -> ExtractedMaterial:
    with path.open("r", encoding="utf-8") as file:
        payload = json.load(file)
    segments = [
        MaterialSegment(location=location, text=text)
        for location, text in _flatten_json(payload)
        if text != "" or "." in location
    ]
    return ExtractedMaterial(
        material_name=path.name,
        material_type=material_type or guess_material_type(path.name),
        file_kind="json",
        segments=segments,
    )


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == value.microsecond == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_section_name(name: str) -> str:
    cleaned = name.split(".", 1)[1].strip() if "." in name else name.strip()
    return "关联交易信息" if cleaned == "关联交易事项" else cleaned


def _is_section_name(name: str) -> bool:
    stripped = name.strip()
    return len(stripped) > 2 and stripped[0].isdigit() and "." in stripped[:3]


def _excel_cell_value(workbook: Any, sheet_name: str, column: str, row: Any) -> str:
    if sheet_name not in workbook.sheetnames:
        return ""
    try:
        row_number = int(row)
    except (TypeError, ValueError):
        return ""
    return _clean_cell(workbook[sheet_name][f"{column}{row_number}"].value)


def parse_excel_template(path: Path, material_type: Optional[str] = None) -> ExtractedMaterial:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=True, keep_vba=False)
    segments: list[MaterialSegment] = []
    try:
        if "要素表" not in workbook.sheetnames:
            for worksheet in workbook.worksheets:
                for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                    values = [_clean_cell(item) for item in row if _clean_cell(item)]
                    if values:
                        segments.append(
                            MaterialSegment(
                                location=f"{worksheet.title}!R{row_number}",
                                text=" | ".join(values),
                            )
                        )
            return ExtractedMaterial(
                material_name=path.name,
                material_type=material_type or guess_material_type(path.name),
                file_kind="excel",
                segments=segments,
            )

        metadata_sheet = workbook["要素表"]
        current_section = ""
        grid_fields: dict[str, dict[str, Any]] = {}
        for row in metadata_sheet.iter_rows(min_row=2, values_only=True):
            values = list(row) + [None] * 18
            name = _clean_cell(values[1])
            value_column = _clean_cell(values[5])
            value_row = _clean_cell(values[6])
            sheet_name = _clean_cell(values[13])
            if not name:
                continue
            if _is_section_name(name):
                current_section = _clean_section_name(name)
                continue
            if not current_section or not value_column or not value_row:
                continue
            if sheet_name and sheet_name != "产品要素":
                grid = grid_fields.setdefault(current_section, {"sheet_name": sheet_name, "fields": []})
                grid["fields"].append((name, value_column))
                continue
            segments.append(
                MaterialSegment(
                    location=f"{current_section}.{name}",
                    text=_excel_cell_value(workbook, "产品要素", value_column, value_row),
                )
            )

        for section, grid in grid_fields.items():
            sheet_name = grid["sheet_name"]
            fields = grid["fields"]
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            row_index = 0
            empty_streak = 0
            for row_number in range(5, worksheet.max_row + 1):
                row_values = [_clean_cell(worksheet[f"{column}{row_number}"].value) for _, column in fields]
                if not any(row_values):
                    empty_streak += 1
                    if empty_streak >= 20 and row_index > 0:
                        break
                    continue
                empty_streak = 0
                for field_index, (field_name, _) in enumerate(fields):
                    segments.append(
                        MaterialSegment(
                            location=f"{section}[{row_index}].{field_name}",
                            text=row_values[field_index],
                        )
                    )
                row_index += 1
    finally:
        workbook.close()

    return ExtractedMaterial(
        material_name=path.name,
        material_type=material_type or guess_material_type(path.name),
        file_kind="excel",
        segments=segments,
    )


def parse_pdf(path: Path, material_type: Optional[str] = None) -> ExtractedMaterial:
    import fitz

    segments: list[MaterialSegment] = []
    with fitz.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            text = (page.get_text("text") or "").strip()
            if not text:
                text = "[本页无可提取文本，疑似扫描件或图片页，当前版本不支持OCR]"
            segments.append(MaterialSegment(location=f"第{page_number}页", text=text))
    return ExtractedMaterial(
        material_name=path.name,
        material_type=material_type or guess_material_type(path.name),
        file_kind="pdf",
        segments=segments,
    )


def parse_docx(path: Path, material_type: Optional[str] = None) -> ExtractedMaterial:
    from docx import Document

    document = Document(str(path))
    segments: list[MaterialSegment] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = (paragraph.text or "").strip()
        if text:
            segments.append(MaterialSegment(location=f"第{index}段", text=text))
    for table_index, table in enumerate(document.tables, start=1):
        rows = [" | ".join((cell.text or "").strip().replace("\n", " ") for cell in row.cells) for row in table.rows]
        if rows:
            segments.append(MaterialSegment(location=f"表{table_index}", text="\n".join(rows)))
    return ExtractedMaterial(
        material_name=path.name,
        material_type=material_type or guess_material_type(path.name),
        file_kind="docx",
        segments=segments,
    )


def parse_txt(path: Path, material_type: Optional[str] = None) -> ExtractedMaterial:
    chunks = path.read_text(encoding="utf-8", errors="ignore").split("\n\n")
    segments = [
        MaterialSegment(location=f"第{index}段", text=chunk.strip())
        for index, chunk in enumerate(chunks, start=1)
        if chunk.strip()
    ]
    return ExtractedMaterial(
        material_name=path.name,
        material_type=material_type or guess_material_type(path.name),
        file_kind="txt",
        segments=segments,
    )


def parse_material(path: Path, material_type: Optional[str] = None) -> ExtractedMaterial:
    if not path.exists():
        raise FileNotFoundError(f"材料文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix == ".json":
        return parse_json_template(path, material_type)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_excel_template(path, material_type)
    if suffix == ".pdf":
        return parse_pdf(path, material_type)
    if suffix == ".docx":
        return parse_docx(path, material_type)
    if suffix == ".txt":
        return parse_txt(path, material_type)
    if suffix == ".doc":
        raise ValueError("不支持老版本 .doc 格式，请另存为 .docx 后上传")
    raise ValueError(f"暂不支持的材料类型：{suffix or '无扩展名'}")
