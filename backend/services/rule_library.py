"""Independent rule-library storage and Excel import support.

This module deliberately does not participate in the review engine. It owns a
separate JSON store used only by the rule-library management page.
"""
from __future__ import annotations

import io
import json
import os
import tempfile
import threading
import time
from pathlib import Path
from typing import Any, Literal, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field, model_validator


LibraryProcess = Literal["pre_registration", "pre_report", "initial", "termination"]
ReviewMethod = Literal["script", "ai"]
RiskLevel = Literal["高风险", "中风险", "低风险"]
RuleObject = Literal["文件", "材料", "表", "字段", "跨材料"]
ReviewDimension = Literal[
    "文件格式标准化",
    "法定要素完整性",
    "文本语义合规",
    "跨文件数据一致性",
    "非标资产穿透",
    "报送时效合规",
]

PROCESS_LABELS: dict[str, str] = {
    "pre_registration": "预登记",
    "pre_report": "事前报告",
    "initial": "初始登记",
    "termination": "终止登记",
}
PROCESS_BY_LABEL = {label: code for code, label in PROCESS_LABELS.items()}

REVIEW_DIMENSIONS = [
    "文件格式标准化",
    "法定要素完整性",
    "文本语义合规",
    "跨文件数据一致性",
    "非标资产穿透",
    "报送时效合规",
]
RULE_OBJECTS = ["文件", "材料", "表", "字段", "跨材料"]
RISK_LEVELS = ["高风险", "中风险", "低风险"]

SUPPORTED_OPERATORS = [
    "enum",
    "dependent_enum",
    "required",
    "max_length",
    "number_precision",
    "date_format",
    "date_range",
    "regex_match",
    "sequence_contiguous",
    "compare_fields",
    "cross_template_equal",
    "conditional_compare",
    "material_required",
    "conditional_material_required",
]

OPERATOR_DESCRIPTIONS: dict[str, str] = {
    "required": "字段必填",
    "enum": "枚举值校验",
    "dependent_enum": "联动枚举校验",
    "max_length": "最大长度校验",
    "number_precision": "数值精度校验",
    "date_format": "日期格式校验",
    "date_range": "日期范围校验",
    "regex_match": "正则格式校验",
    "sequence_contiguous": "序号连续性校验",
    "compare_fields": "字段数值比较",
    "cross_template_equal": "跨模板字段一致性",
    "conditional_compare": "条件触发校验",
    "material_required": "材料必交校验",
    "conditional_material_required": "条件性材料必交",
}

RULE_LIBRARY_FILE = Path(
    os.getenv(
        "RULE_LIBRARY_FILE",
        Path(__file__).resolve().parent.parent / "data" / "rule_library" / "rules.json",
    )
)
_STORE_LOCK = threading.RLock()


class LibraryRuleInput(BaseModel):
    rule_id: str = Field(min_length=3, max_length=80, pattern=r"^[A-Za-z0-9_-]+$")
    rule_name: str = Field(min_length=1, max_length=200)
    process: LibraryProcess
    review_method: ReviewMethod
    applicable_materials: list[str] = Field(default_factory=list)
    rule_object: RuleObject
    table_name: str = ""
    field_path: str = ""
    review_dimension: ReviewDimension
    trigger_condition: str = ""
    rule_text: str = Field(min_length=1)
    basis_text: str = Field(min_length=1)
    risk_level: RiskLevel = "中风险"
    version: str = "V1.0"
    enabled: bool = True
    remarks: str = ""
    operator: str = ""
    script_params: dict[str, Any] = Field(default_factory=dict)
    special_prompt: str = ""

    @model_validator(mode="after")
    def validate_rule_shape(self):
        self.rule_name = self.rule_name.strip()
        self.table_name = self.table_name.strip()
        self.field_path = self.field_path.strip()
        self.trigger_condition = self.trigger_condition.strip()
        self.rule_text = self.rule_text.strip()
        self.basis_text = self.basis_text.strip()
        self.version = self.version.strip() or "V1.0"
        self.remarks = self.remarks.strip()
        self.operator = self.operator.strip()
        self.special_prompt = self.special_prompt.strip()
        self.applicable_materials = [
            str(item).strip() for item in self.applicable_materials if str(item).strip()
        ]

        if not self.applicable_materials:
            raise ValueError("适用材料至少填写一项")
        if self.rule_object == "字段" and not self.field_path:
            raise ValueError("字段类规则必须填写字段路径")
        if self.rule_object == "表" and not self.table_name:
            raise ValueError("表类规则必须填写表名")
        if self.review_method == "script":
            if self.operator not in SUPPORTED_OPERATORS:
                raise ValueError("脚本规则必须选择受支持的 operator")
            self.special_prompt = ""
        else:
            if not self.special_prompt:
                raise ValueError("AI规则必须填写专属提示词")
            self.operator = ""
            self.script_params = {}
        return self


class LibraryRule(LibraryRuleInput):
    source_file: str = ""
    source_sheet: str = ""
    source_row: int = 0
    created_at: int
    updated_at: int


class ImportIssue(BaseModel):
    sheet: str = ""
    row: int = 0
    rule_id: str = ""
    message: str


class ImportReport(BaseModel):
    filename: str
    process: LibraryProcess
    total_rules: int = 0
    script_count: int = 0
    ai_count: int = 0
    valid: bool = False
    errors: list[ImportIssue] = Field(default_factory=list)
    rules: list[LibraryRule] = Field(default_factory=list)


def _load_all() -> list[LibraryRule]:
    with _STORE_LOCK:
        if not RULE_LIBRARY_FILE.exists():
            return []
        raw = json.loads(RULE_LIBRARY_FILE.read_text(encoding="utf-8"))
        return [LibraryRule.model_validate(item) for item in raw.get("rules", [])]


def _save_all(rules: list[LibraryRule]) -> None:
    with _STORE_LOCK:
        RULE_LIBRARY_FILE.parent.mkdir(parents=True, exist_ok=True)
        payload = json.dumps(
            {"rules": [rule.model_dump() for rule in rules]},
            ensure_ascii=False,
            indent=2,
        )
        fd, tmp_name = tempfile.mkstemp(
            prefix="rules-", suffix=".tmp", dir=str(RULE_LIBRARY_FILE.parent)
        )
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as handle:
                handle.write(payload)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(tmp_name, RULE_LIBRARY_FILE)
        finally:
            if os.path.exists(tmp_name):
                os.unlink(tmp_name)


def list_rules(process: Optional[LibraryProcess] = None) -> list[LibraryRule]:
    rules = _load_all()
    if process:
        rules = [rule for rule in rules if rule.process == process]
    return sorted(rules, key=lambda item: (item.updated_at, item.rule_id), reverse=True)


def get_rule(rule_id: str) -> Optional[LibraryRule]:
    return next((rule for rule in _load_all() if rule.rule_id == rule_id), None)


def create_rule(payload: LibraryRuleInput, *, source: dict[str, Any] | None = None) -> LibraryRule:
    rules = _load_all()
    if any(rule.rule_id == payload.rule_id for rule in rules):
        raise ValueError(f"规则ID {payload.rule_id} 已存在")
    now = int(time.time())
    source = source or {}
    rule = LibraryRule(
        **payload.model_dump(),
        source_file=str(source.get("source_file") or ""),
        source_sheet=str(source.get("source_sheet") or ""),
        source_row=int(source.get("source_row") or 0),
        created_at=now,
        updated_at=now,
    )
    rules.append(rule)
    _save_all(rules)
    return rule


def create_rules_bulk(items: list[LibraryRule]) -> list[LibraryRule]:
    if not items:
        return []
    rules = _load_all()
    existing_ids = {rule.rule_id for rule in rules}
    incoming_ids = [rule.rule_id for rule in items]
    duplicates = sorted({rule_id for rule_id in incoming_ids if rule_id in existing_ids})
    if len(incoming_ids) != len(set(incoming_ids)):
        duplicates.extend(
            sorted({rule_id for rule_id in incoming_ids if incoming_ids.count(rule_id) > 1})
        )
    if duplicates:
        raise ValueError(f"规则ID重复：{', '.join(sorted(set(duplicates)))}")
    rules.extend(items)
    _save_all(rules)
    return items


def update_rule(rule_id: str, payload: LibraryRuleInput) -> Optional[LibraryRule]:
    rules = _load_all()
    current_index = next((idx for idx, rule in enumerate(rules) if rule.rule_id == rule_id), None)
    if current_index is None:
        return None
    if payload.rule_id != rule_id and any(rule.rule_id == payload.rule_id for rule in rules):
        raise ValueError(f"规则ID {payload.rule_id} 已存在")
    current = rules[current_index]
    updated = LibraryRule(
        **payload.model_dump(),
        source_file=current.source_file,
        source_sheet=current.source_sheet,
        source_row=current.source_row,
        created_at=current.created_at,
        updated_at=int(time.time()),
    )
    rules[current_index] = updated
    _save_all(rules)
    return updated


def set_rule_enabled(rule_id: str, enabled: bool) -> Optional[LibraryRule]:
    current = get_rule(rule_id)
    if current is None:
        return None
    payload = LibraryRuleInput.model_validate({**current.model_dump(), "enabled": enabled})
    return update_rule(rule_id, payload)


def copy_rule(rule_id: str) -> Optional[LibraryRule]:
    current = get_rule(rule_id)
    if current is None:
        return None
    existing = {rule.rule_id for rule in _load_all()}
    base = f"{current.rule_id}-COPY"
    candidate = base
    sequence = 2
    while candidate in existing:
        candidate = f"{base}-{sequence}"
        sequence += 1
    payload = LibraryRuleInput.model_validate(
        {
            **current.model_dump(),
            "rule_id": candidate,
            "rule_name": f"{current.rule_name}（副本）",
            "enabled": False,
        }
    )
    return create_rule(
        payload,
        source={
            "source_file": current.source_file,
            "source_sheet": current.source_sheet,
            "source_row": current.source_row,
        },
    )


def delete_rule(rule_id: str) -> bool:
    rules = _load_all()
    remaining = [rule for rule in rules if rule.rule_id != rule_id]
    if len(remaining) == len(rules):
        return False
    _save_all(remaining)
    return True


def delete_rules(rule_ids: list[str]) -> int:
    target = set(rule_ids)
    rules = _load_all()
    remaining = [rule for rule in rules if rule.rule_id not in target]
    deleted = len(rules) - len(remaining)
    if deleted:
        _save_all(remaining)
    return deleted


def build_visual_rule(rule: LibraryRuleInput | LibraryRule) -> str:
    if rule.review_method != "script":
        return ""
    field = rule.field_path or rule.table_name or rule.rule_name
    params = rule.script_params or {}
    operator = rule.operator
    detail = OPERATOR_DESCRIPTIONS.get(operator, operator)
    if operator == "required":
        detail = "不得为空"
    elif operator == "enum":
        detail = f"只能填写：{'、'.join(map(str, params.get('values') or [])) or '规定枚举值'}"
    elif operator == "max_length":
        detail = f"长度不得超过 {params.get('max_length', '')} 个字符"
    elif operator == "number_precision":
        detail = f"最多保留 {params.get('scale', 2)} 位小数"
    elif operator == "date_format":
        detail = "必须符合 YYYY-MM-DD 日期格式"
    elif operator == "date_range":
        detail = f"日期范围：{params.get('min', 'today')} 至未来 {params.get('max_days', '')} 天"
    elif operator == "regex_match":
        detail = f"必须匹配格式 {params.get('pattern', '')}"
    elif operator == "compare_fields":
        detail = f"不得大于字段 {params.get('right_field', '')}"
    elif operator == "cross_template_equal":
        detail = f"必须与基准字段 {params.get('baseline_field', field)} 一致"
    elif operator in {"material_required", "conditional_material_required"}:
        detail = f"必须提交材料 {params.get('material_label', rule.rule_name)}"
    elif operator == "conditional_compare":
        trigger = params.get("trigger") or {}
        target = params.get("target") or {}
        detail = (
            f"当 {trigger.get('field', '')} {trigger.get('op', '等于')} {trigger.get('value', '')} 时，"
            f"{target.get('field', field)} 执行 {target.get('op', '校验')}"
        )
    prefix = f"当 {rule.trigger_condition} 时，" if rule.trigger_condition else ""
    return f"{prefix}{field} {detail}".strip()


COMMON_HEADERS = [
    "规则ID",
    "规则名称",
    "适用流程",
    "审核方式",
    "适用材料",
    "规则对象",
    "要素分类/表名",
    "字段路径",
    "审查维度",
    "触发条件",
    "具体规则",
    "审查依据",
    "风险等级",
    "当前版本",
    "启用状态",
    "备注",
]
SCRIPT_HEADERS = COMMON_HEADERS + ["operator", "参数(JSON)"]
AI_HEADERS = COMMON_HEADERS + ["专属提示词"]

COLUMN_DESCRIPTIONS: dict[str, str] = {
    "规则ID": "全局唯一，仅支持3-80位字母、数字、下划线和连字符。",
    "规则名称": "便于列表检索和详情展示的业务名称。",
    "适用流程": "必须与上传前选择的登记流程一致。",
    "审核方式": "脚本审核规则 Sheet 填脚本规则，AI审核规则 Sheet 填 AI规则。",
    "适用材料": "多个材料用顿号、逗号或换行分隔。",
    "规则对象": "从文件、材料、表、字段、跨材料中选择。",
    "要素分类/表名": "表类规则必填；字段类规则可填写所属表名或要素分类。",
    "字段路径": "字段类规则必填，例如：产品基本信息.募集规模。",
    "审查维度": "从六个固定审查维度中选择。",
    "触发条件": "可选，填写规则生效的业务前置条件。",
    "具体规则": "正式业务审核要求，必须填写。",
    "审查依据": "法规、业务指南、填表说明或内部审核口径，必须填写。",
    "风险等级": "从高风险、中风险、低风险中选择。",
    "当前版本": "普通文本版本号，新增规则建议使用 V1.0。",
    "启用状态": "从启用或停用中选择。",
    "备注": "可选补充说明。",
    "operator": "脚本规则的结构化操作符，详见 operator字典 Sheet。",
    "参数(JSON)": "脚本规则参数对象，必须是合法 JSON。",
    "专属提示词": "AI执行说明，包含判定重点、排除条件和证据要求。",
}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _split_materials(value: str) -> list[str]:
    if not value:
        return []
    import re

    return [item.strip() for item in re.split(r"[、,，;；/\n]+", value) if item.strip()]


def _parse_enabled(value: str) -> bool:
    if value in {"启用", "是", "true", "True", "1"}:
        return True
    if value in {"停用", "否", "false", "False", "0"}:
        return False
    raise ValueError("启用状态只能填写启用或停用")


def _parse_process(value: str) -> str:
    return PROCESS_BY_LABEL.get(value, value)


def _parse_review_method(value: str) -> str:
    if value in {"脚本规则", "脚本", "script"}:
        return "script"
    if value in {"AI规则", "AI", "ai"}:
        return "ai"
    return value


def parse_rule_workbook(
    content: bytes,
    *,
    filename: str,
    process: LibraryProcess,
    existing_ids: set[str] | None = None,
) -> ImportReport:
    existing_ids = existing_ids or set()
    errors: list[ImportIssue] = []
    rules: list[LibraryRule] = []
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        return ImportReport(
            filename=filename,
            process=process,
            errors=[ImportIssue(message=f"Excel文件无法读取：{exc}")],
        )

    available = [
        ("脚本审核规则", "script", SCRIPT_HEADERS),
        ("AI审核规则", "ai", AI_HEADERS),
    ]
    found_sheet = False
    seen_ids: set[str] = set()
    now = int(time.time())
    try:
        for sheet_name, method, expected_headers in available:
            if sheet_name not in workbook.sheetnames:
                continue
            found_sheet = True
            sheet = workbook[sheet_name]
            mapping = {
                _cell_text(cell.value): index for index, cell in enumerate(sheet[1]) if _cell_text(cell.value)
            }
            missing_headers = [header for header in expected_headers if header not in mapping]
            if missing_headers:
                errors.append(
                    ImportIssue(sheet=sheet_name, message=f"缺少列：{'、'.join(missing_headers)}")
                )
                continue

            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if not any(_cell_text(value) for value in row):
                    continue

                def value(header: str) -> str:
                    index = mapping.get(header)
                    return _cell_text(row[index]) if index is not None and index < len(row) else ""

                meaningful_headers = [
                    header for header in expected_headers if header not in {"适用流程", "审核方式"}
                ]
                if not any(value(header) for header in meaningful_headers):
                    continue

                rule_id = value("规则ID")
                try:
                    row_process = _parse_process(value("适用流程"))
                    row_method = _parse_review_method(value("审核方式"))
                    if row_process != process:
                        raise ValueError(
                            f"适用流程必须为当前流程 {PROCESS_LABELS[process]}"
                        )
                    if row_method != method:
                        expected = "脚本规则" if method == "script" else "AI规则"
                        raise ValueError(f"审核方式必须为 {expected}")
                    if rule_id in seen_ids:
                        raise ValueError("规则ID在当前Excel中重复")
                    if rule_id in existing_ids:
                        raise ValueError("规则ID已存在于规则库")
                    params: dict[str, Any] = {}
                    if method == "script":
                        raw_params = value("参数(JSON)") or "{}"
                        try:
                            params = json.loads(raw_params)
                        except json.JSONDecodeError as exc:
                            raise ValueError(f"参数(JSON)格式错误：{exc.msg}") from exc
                        if not isinstance(params, dict):
                            raise ValueError("参数(JSON)必须是对象")
                    payload = LibraryRuleInput(
                        rule_id=rule_id,
                        rule_name=value("规则名称"),
                        process=process,
                        review_method=method,
                        applicable_materials=_split_materials(value("适用材料")),
                        rule_object=value("规则对象"),
                        table_name=value("要素分类/表名"),
                        field_path=value("字段路径"),
                        review_dimension=value("审查维度"),
                        trigger_condition=value("触发条件"),
                        rule_text=value("具体规则"),
                        basis_text=value("审查依据"),
                        risk_level=value("风险等级"),
                        version=value("当前版本"),
                        enabled=_parse_enabled(value("启用状态")),
                        remarks=value("备注"),
                        operator=value("operator") if method == "script" else "",
                        script_params=params,
                        special_prompt=value("专属提示词") if method == "ai" else "",
                    )
                    seen_ids.add(rule_id)
                    rules.append(
                        LibraryRule(
                            **payload.model_dump(),
                            source_file=filename,
                            source_sheet=sheet_name,
                            source_row=row_number,
                            created_at=now,
                            updated_at=now,
                        )
                    )
                except Exception as exc:
                    errors.append(
                        ImportIssue(
                            sheet=sheet_name,
                            row=row_number,
                            rule_id=rule_id,
                            message=str(exc),
                        )
                    )
    finally:
        workbook.close()

    if not found_sheet:
        errors.append(ImportIssue(message="至少需要脚本审核规则或AI审核规则Sheet"))
    if found_sheet and not rules and not errors:
        errors.append(ImportIssue(message="未识别到可导入规则"))

    return ImportReport(
        filename=filename,
        process=process,
        total_rules=len(rules),
        script_count=sum(rule.review_method == "script" for rule in rules),
        ai_count=sum(rule.review_method == "ai" for rule in rules),
        valid=not errors and bool(rules),
        errors=errors,
        rules=rules,
    )


def build_import_template(process: LibraryProcess) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="4A54A8")
    header_font = Font(color="FFFFFF", bold=True)

    def add_sheet(name: str, headers: list[str], method_label: str):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.comment = Comment(COLUMN_DESCRIPTIONS.get(str(cell.value), ""), "规则库")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headers)).coordinate}"
        for index, header in enumerate(headers, start=1):
            width = 18
            if header in {"具体规则", "审查依据", "专属提示词", "参数(JSON)"}:
                width = 32
            elif header in {"适用材料", "字段路径", "触发条件"}:
                width = 24
            sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width

        process_col = headers.index("适用流程") + 1
        method_col = headers.index("审核方式") + 1
        sheet.cell(2, process_col, PROCESS_LABELS[process])
        sheet.cell(2, method_col, method_label)

        validations = {
            "适用流程": list(PROCESS_LABELS.values()),
            "审核方式": [method_label],
            "规则对象": RULE_OBJECTS,
            "审查维度": REVIEW_DIMENSIONS,
            "风险等级": RISK_LEVELS,
            "启用状态": ["启用", "停用"],
        }
        if "operator" in headers:
            validations["operator"] = SUPPORTED_OPERATORS
        for header, values in validations.items():
            column = headers.index(header) + 1
            formula = '"' + ",".join(values) + '"'
            validation = DataValidation(type="list", formula1=formula, allow_blank=False)
            sheet.add_data_validation(validation)
            validation.add(f"{sheet.cell(2, column).coordinate}:{sheet.cell(201, column).coordinate}")

    add_sheet("脚本审核规则", SCRIPT_HEADERS, "脚本规则")
    add_sheet("AI审核规则", AI_HEADERS, "AI规则")

    descriptions = workbook.create_sheet("列说明")
    descriptions.append(["字段", "填写说明"])
    for cell in descriptions[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for header in COMMON_HEADERS + ["operator", "参数(JSON)", "专属提示词"]:
        descriptions.append([header, COLUMN_DESCRIPTIONS[header]])
    descriptions.column_dimensions["A"].width = 24
    descriptions.column_dimensions["B"].width = 88
    descriptions.freeze_panes = "A2"

    dictionary = workbook.create_sheet("operator字典")
    dictionary.append(["operator", "中文说明", "参数示例"])
    examples = {
        "required": "{}",
        "enum": '{"values":["是","否"]}',
        "dependent_enum": '{"parent_field":"A","child_field":"B","mapping":{"A1":["B1"]}}',
        "max_length": '{"max_length":50}',
        "number_precision": '{"scale":2}',
        "date_format": "{}",
        "date_range": '{"min":"today","max_days":365}',
        "regex_match": '{"pattern":"[A-Z]{2}\\\\d{6}"}',
        "sequence_contiguous": "{}",
        "compare_fields": '{"right_field":"发行规模"}',
        "cross_template_equal": '{"current_field":"产品名称","baseline_field":"产品名称"}',
        "conditional_compare": '{"trigger":{"field":"A","op":"equals","value":"是"},"target":{"field":"B","op":"required"}}',
        "material_required": '{"material_label":"信托合同","file_ext":[".pdf"]}',
        "conditional_material_required": '{"material_label":"关联交易说明"}',
    }
    for cell in dictionary[1]:
        cell.fill = header_fill
        cell.font = header_font
    for operator in SUPPORTED_OPERATORS:
        dictionary.append([operator, OPERATOR_DESCRIPTIONS[operator], examples[operator]])
    dictionary.column_dimensions["A"].width = 34
    dictionary.column_dimensions["B"].width = 28
    dictionary.column_dimensions["C"].width = 74
    dictionary.freeze_panes = "A2"

    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()
